"""Supplier identities and full-delivery handback use real DB and HTTP APIs."""
from io import BytesIO
from uuid import UUID,uuid4
import pytest
from openpyxl import Workbook,load_workbook
from sqlalchemy import text
from test_auth_postgres import env,client,login
from test_business_registry_postgres import allocated
from test_documents_postgres import body as invoice
from marking_erp.business import store
from marking_erp.business.worker import background

pytestmark=[pytest.mark.postgres,pytest.mark.usefixtures('fresh_schema')]


def full_excel(quantity=10):
    book=Workbook();sheet=book.active
    for row in [['PN','GTIN','Количество'],['PN/A-1','04640638349704',quantity],['BOLT','',5],['BOLT','',7]]:sheet.append(row)
    result=BytesIO();book.save(result);return result.getvalue()


def request_full(env,engine,item,code):
    # Local upstream-fact fixture only; no real external success is manufactured.
    with engine.begin() as c:
        actor=background(c,env[3]['specialist'],env[1])
        c.execute(text("UPDATE business_items SET state='applied',body=body-'warehouse_composition' WHERE id=:id"),{'id':item})
        c.execute(text("UPDATE business_codes SET state='applied' WHERE id=:id"),{'id':code['id']})
        store.create(actor,'atk','ATK fixture',{'supplier':'Поставщик А','tnved':'8542399000','documents':[str(item)],'codes':[str(code['id'])],'atk':'1'*25},env[1],state='transferred')
    with client(env) as c:
        login(c,'logistics');doc=c.get('/api/logistics/documents/'+str(item)).json()['document']
        result=c.post('/api/logistics/documents/'+str(item)+'/request-composition',json={'revision':doc['revision'],'confirmation':'INV-01'})
        assert result.status_code==200,result.text


def test_logistics_returns_to_bm_full_excel_adds_only_information_and_keeps_km(allocated,engine):
    env,item,source,code=allocated;request_full(env,engine,item,code);data=full_excel()
    with client(env) as c:
        login(c);doc=c.get('/api/documents/'+str(source)).json()['document']
        assert doc['warehouse_composition']['state']=='requested'
        queue=c.get('/api/documents',params={'organization_id':str(env[1]),'stage':'composition_requested'}).json()['documents']
        assert [v['id'] for v in queue]==[str(source)]
        params={'revision':doc['marking_revision'],'source_revision':doc['revision']}
        preview=c.post('/api/documents/'+str(source)+'/warehouse-composition/preview',params=params,content=data)
        assert preview.status_code==200,preview.text
        parsed=preview.json();assert not parsed['errors'] and len(parsed['lines'])==2
        assert [v['quantity'] for v in parsed['lines']]==[5,7]
        saved=c.post('/api/documents/'+str(source)+'/warehouse-composition/commit',params={**params,'sha256':parsed['sha256']},content=data)
        assert saved.status_code==200,saved.text
        repeated=c.post('/api/documents/'+str(source)+'/warehouse-composition/commit',params={**params,'sha256':parsed['sha256']},content=data)
        assert repeated.status_code==200 and repeated.json()['replayed']
        doc=c.get('/api/documents/'+str(source)).json()['document'];assert doc['revision']==3
        assert doc['lines'][0]['quantity']==10 and doc['lines'][0]['labels']==1
        assert [v['marked'] for v in doc['lines'][1:]]==[False,False]
        assert c.get('/api/business-files/'+doc['warehouse_composition']['file_id']).content==data
    with engine.connect() as c:
        actual=c.execute(text('SELECT * FROM business_codes WHERE id=:id'),{'id':code['id']}).mappings().one()
        assert bytes(actual['payload'])==bytes(code['payload']) and actual['registry_payload_id']==code['registry_payload_id']
        assert c.execute(text('SELECT count(*) FROM external_operations')).scalar_one()==0
    with client(env) as c:
        login(c,'logistics');task=uuid4()
        assert c.get('/api/business-files/'+doc['warehouse_composition']['file_id']).content==data
        result=c.post('/api/logistics/warehouse-transfer',json={'request_id':str(task),'source_id':str(source),'source_revision':3,'confirmed':True,'kind':'shipment','assignees':[str(env[3]['warehouse'])]});assert result.status_code==200,result.text
    with client(env) as c:
        login(c,'warehouse');detail=c.get('/api/warehouse/tasks/'+str(task)).json();assert len(detail['lines'])==3
        result=c.post('/api/warehouse/tasks/'+str(task)+'/scans',json={'request_id':str(uuid4()),'raw':bytes(code['payload']).decode()});assert result.status_code==200,result.text
        detail=c.get('/api/warehouse/tasks/'+str(task)).json()
        assert c.post('/api/warehouse/tasks/'+str(task)+'/close',json={'revision':detail['task']['revision'],'confirmation':'INV-01'}).status_code==200
        export=c.get('/api/warehouse/tasks/'+str(task)+'/export');book=load_workbook(BytesIO(export.content))
        values=list(book['Немаркированные строки'].values);assert [r[1] for r in values[1:]]==[5,7]


def test_full_invoice_rejects_marked_changes_and_wrong_role_and_file(allocated,engine):
    env,item,source,code=allocated
    with engine.begin() as db:db.execute(text("UPDATE business_items SET body=body-'warehouse_composition' WHERE id=:id"),{'id':item})
    with client(env) as c:
        login(c,'logistics')
        assert c.post('/api/logistics/warehouse-transfer',json={'request_id':str(uuid4()),'source_id':str(source),'source_revision':2,'confirmed':True,'kind':'receipt','assignees':[str(env[3]['warehouse'])]}).status_code==409
    request_full(env,engine,item,code)
    with client(env) as c:
        login(c);doc=c.get('/api/documents/'+str(source)).json()['document'];params={'revision':doc['marking_revision'],'source_revision':doc['revision']}
        wrong=c.post('/api/documents/'+str(source)+'/warehouse-composition/preview',params=params,content=full_excel(11)).json()
        assert any('Изменять маркированную часть нельзя' in e['message'] for e in wrong['errors'])
        correct=c.post('/api/documents/'+str(source)+'/warehouse-composition/preview',params=params,content=full_excel()).json()
        assert c.post('/api/documents/'+str(source)+'/warehouse-composition/commit',params={**params,'sha256':correct['sha256']},content=full_excel(11)).status_code==409
        assert c.get('/api/documents/'+str(source)).json()['document']['revision']==2
        with engine.begin() as db:db.execute(text("UPDATE group_grants SET scope='none' WHERE group_slug='bm' AND permission_id='marking.documents.view'"))
        assert c.get('/api/documents/'+str(source)+'/warehouse-composition/template').status_code==403
        assert c.post('/api/documents/'+str(source)+'/warehouse-composition/preview',params=params,content=full_excel()).status_code==403
    with client(env) as c:
        login(c,'warehouse')
        assert c.post('/api/documents/'+str(source)+'/warehouse-composition/preview',params=params,content=full_excel()).status_code==403


def test_supplier_merge_flattening_stale_revision_and_source_history(env,engine):
    with client(env) as c:
        login(c);sources=[]
        for n,name in enumerate(['ООО Поставщик','Поставщик ООО','Поставщик-3']):
            response=c.post('/api/documents',json=invoice(env,number='S-'+str(n),supplier=name));assert response.status_code==200,response.text;sources.append(response.json()['id'])
    with client(env) as c:
        login(c,'administrator');data={'organization_id':str(env[1]),'revision':0,'aliases':['ООО Поставщик','Поставщик ООО'],'canonical':'Поставщик','reason':'One supplier spelling variants'}
        assert c.post('/api/suppliers/merge',json=data).status_code==200
        assert c.post('/api/suppliers/merge',json=data).status_code==409
        assert c.post('/api/suppliers/merge',json={**data,'revision':1,'aliases':['Поставщик','Поставщик-3'],'canonical':'Поставщик ООО'}).status_code==200
        values=c.get('/api/suppliers',params={'organization_id':str(env[1])}).json()['suppliers']
        assert {v['canonical'] for v in values}=={'Поставщик'}
        assert c.post('/api/suppliers/separate',json={**data,'revision':2,'aliases':['Поставщик-3']}).status_code==200
    with client(env) as c:
        login(c);doc=c.get('/api/documents/'+sources[0]).json()['document']
        assert doc['supplier']=='ООО Поставщик' and doc['canonical_supplier']=='Поставщик'
        assert c.get('/api/suppliers',params={'organization_id':str(env[1])}).status_code==403
        duplicate=c.post('/api/documents',json=invoice(env,number='S-0',supplier='Поставщик ООО'));assert duplicate.status_code==409,duplicate.text
        assert c.get('/api/documents',params={'organization_id':str(env[1]),'q':'Поставщик'}).json()['documents']
    with engine.connect() as c:
        assert c.execute(text("SELECT snapshot->>'supplier' FROM bm_document_versions WHERE document_id=:id AND revision=1"),{'id':sources[0]}).scalar_one()=='ООО Поставщик'


def test_supplier_aliases_unify_pending_batch_without_allowing_subset(env,engine):
    with client(env) as c:
        login(c)
        for n,name in enumerate(['NAME-A','NAME-B']):assert c.post('/api/documents',json=invoice(env,number='BATCH-'+str(n),supplier=name)).status_code==200
    with client(env) as c:
        login(c,'administrator');assert c.post('/api/suppliers/merge',json={'organization_id':str(env[1]),'revision':0,'aliases':['NAME-A','NAME-B'],'canonical':'CANONICAL','reason':'Same supplier aliases'}).status_code==200
    with client(env) as c:
        login(c);suppliers=c.get('/api/atks/supplier-batches',params={'organization_id':str(env[1])}).json()
        assert suppliers['suppliers']==['CANONICAL']
        batch=c.get('/api/atks/supplier-batch',params={'organization_id':str(env[1]),'supplier':'CANONICAL'}).json();assert len(batch['documents'])==2 and not batch['ready']
