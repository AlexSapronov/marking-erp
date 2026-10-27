from uuid import uuid4
from io import BytesIO
import pytest
from sqlalchemy import text
from openpyxl import Workbook,load_workbook
from test_auth_postgres import env,client,login
from test_documents_postgres import body as invoice
from marking_erp.registry import register_product

pytestmark=[pytest.mark.postgres,pytest.mark.usefixtures('fresh_schema')]
GTIN='04640638349704'

def card(env,**extra):return {'organization_id':str(env[1]),'request_id':str(uuid4()),'pn':'PN-100','gtin':GTIN,'tnved':'8542399000','name':'PN-100 component','brand':'Brand','category_id':1,'attributes':{},**extra}

def connections(c,env,**extra):
    result=c.put('/api/connections',json={'organization_id':str(env[1]),'revision':0,'environment':'production','writes_enabled':False,'warehouse_live_checks':False,'template_id':36,**extra})
    assert result.status_code==200,result.text


def test_card_duplicate_acknowledgement_and_idempotent_create(env):
    with client(env) as c:
        login(c);value=card(env);first=c.post('/api/catalog',json=value);assert first.status_code==200,first.text
        assert c.post('/api/catalog',json=value).json()['id']==first.json()['id']
        duplicate=c.post('/api/catalog',json={**value,'request_id':str(uuid4())});assert duplicate.status_code==409
        assert duplicate.json()['detail']['candidates'][0]['id']==first.json()['id']
        assert c.post('/api/catalog',json={**value,'request_id':str(uuid4()),'duplicate_reason':'Different manufacturer entry'}).status_code==200
    with client(env) as c:
        login(c,'tsd');assert c.get('/api/catalog',params={'organization_id':str(env[1])}).status_code==403


def test_approved_shared_template_applies_to_own_cards(env):
    with client(env) as bm,client(env) as specialist:
        login(bm);login(specialist,'specialist');value=bm.post('/api/catalog',json=card(env)).json()
        template=specialist.post('/api/catalog/templates',json={'organization_id':str(env[1]),'title':'Shared','tnved':'8542399000','category_id':1,'attributes':{'12':'Component'}});assert template.status_code==200,template.text
        response=bm.post('/api/catalog/templates/'+template.json()['id']+'/apply',json={'ids':[value['id']],'revision':1});assert response.status_code==200,response.text
        assert bm.get('/api/catalog/'+value['id']).json()['item']['body']['attributes']['12']=='Component'


def test_xlsx_import_and_export_preserve_text_pn_and_never_allocate_gtin_locally(env,engine):
    book=Workbook();book.active.append(['PN','ТН ВЭД','Товарный знак']);book.active.append(['00123','8542399000','Brand']);out=BytesIO();book.save(out)
    with client(env) as c:
        login(c);preview=c.post('/api/catalog-import/preview',params={'organization_id':str(env[1])},content=out.getvalue());assert preview.status_code==200,preview.text
        id=preview.json()['id'];committed=c.post('/api/catalog-import/'+id+'/commit',json={'ids':[id],'revision':1});assert committed.status_code==200,committed.text
        exported=c.get('/api/catalog/export',params={'organization_id':str(env[1])});assert exported.status_code==200
        rows=list(load_workbook(BytesIO(exported.content)).active.values);assert rows[1][0]=='00123' and rows[1][1] in ('',None)
    with engine.connect() as c:assert c.execute(text('SELECT count(*) FROM external_operations')).scalar_one()==0


def test_integration_secrets_are_encrypted_and_not_echoed(env,engine):
    secret='local-fixture-nk-key'
    with client(env) as c:
        login(c,'administrator');connections(c,env,nk_key=secret)
        response=c.get('/api/connections',params={'organization_id':str(env[1])});assert response.status_code==200
        assert secret not in response.text and response.json()['credentials']['nk_key'] is True
    with engine.connect() as c:
        ciphertext=c.execute(text('SELECT secret FROM integration_settings')).scalar_one();assert secret.encode() not in bytes(ciphertext)
    with client(env) as c:
        login(c,'specialist');assert c.get('/api/connections',params={'organization_id':str(env[1])}).status_code==403


def test_emission_prepares_signature_job_without_claiming_codes_or_success(env,engine):
    with client(env) as admin:
        login(admin,'administrator');connections(admin,env,writes_enabled=True)
    # Synthetic fixture only; no remote NK response or real emission is simulated.
    with engine.begin() as c:register_product(c,env[1],GTIN,'PN/A-1','8542399000',source='external_verified',source_ref='local-test-fixture',evidence_sha256='0'*64)
    with client(env) as c:
        login(c);created=c.post('/api/documents',json=invoice(env));assert created.status_code==200
        id=created.json()['id'];assert c.post('/api/documents/'+id+'/prepare',json={'revision':1}).status_code==200
        body={'source_id':id,'revision':2,'request_id':str(uuid4())};response=c.post('/api/marking/start',json=body);assert response.status_code==200,response.text
        data=c.get('/api/business/'+response.json()['id']).json();assert not data['codes'] and not data['files']
        assert data['operations'][0]['state']=='awaiting_signature'
        assert c.post('/api/marking/start',json=body).json()['id']==response.json()['id']
    with engine.connect() as c:assert c.execute(text('SELECT count(*) FROM business_codes')).scalar_one()==0


def test_user_session_cannot_attest_signature_verification(env):
    with client(env) as c:
        login(c,'administrator');response=c.post('/api/signer/verification',json={'job_id':str(uuid4()),'fence':1,'payload_sha256':'0'*64,'signature_sha256':'0'*64,'certificate_sha256':'0'*64,'cades_valid':True,'chain_valid':True})
        assert response.status_code==401


def test_catalog_export_uses_filter_and_literal_wildcards(env):
    with client(env) as c:
        login(c)
        for pn in ['EXPORT-ONE','EXPORT-TWO']:
            response=c.post('/api/catalog',json=card(env,pn=pn,duplicate_reason='Separate model fixture'));assert response.status_code==200,response.text
        response=c.get('/api/catalog/export',params={'organization_id':str(env[1]),'q':'EXPORT-ONE'})
        values=list(load_workbook(BytesIO(response.content)).active.values)
        assert len(values)==2 and values[1][0]=='EXPORT-ONE'
        response=c.get('/api/catalog/export',params={'organization_id':str(env[1]),'q':'%'})
        assert len(list(load_workbook(BytesIO(response.content)).active.values))==1
