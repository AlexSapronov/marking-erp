"""NK orchestration against PostgreSQL; external reads use explicit contract fixtures."""
import json
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from uuid import UUID, uuid4
import pytest
from openpyxl import Workbook,load_workbook
from sqlalchemy import text
from test_auth_postgres import env, client, login
from test_business_postgres import card, connections
from test_catalog_worklist_postgres import catalog_fixture, listed
from marking_erp.business import store
from marking_erp.business.catalog_schedule import tick
from marking_erp.business.worker import BusinessWorker,background
from marking_erp.business.nk import Client

pytestmark=[pytest.mark.postgres,pytest.mark.usefixtures('fresh_schema')]


def rule(engine,key,value):
    with engine.begin() as c:c.execute(text("INSERT INTO policy_values VALUES ('system',:key,CAST(:value AS jsonb)) ON CONFLICT(scope_key,rule_id) DO UPDATE SET value=EXCLUDED.value"),{'key':key,'value':json.dumps(value)})


def test_effective_duplicate_rules_warn_without_merging(env,engine):
    rule(engine,'nk.duplicates.similarity',1.0)
    with client(env) as c:
        login(c)
        assert c.post('/api/catalog',json=card(env,gtin='',pn='COMPONENT-100')).status_code==200
        assert c.post('/api/catalog',json=card(env,gtin='',pn='COMPONENT-101')).status_code==200
        rule(engine,'nk.duplicates.similarity',0.5)
        proposed=card(env,gtin='',pn='COMPONENT-102')
        assert c.post('/api/catalog',json=proposed).status_code==409
        rule(engine,'nk.duplicates.reason_required',False)
        assert c.post('/api/catalog',json=proposed).status_code==409
        assert c.post('/api/catalog',json={**proposed,'duplicate_reason':'Да'}).status_code==200
    with engine.connect() as c:
        assert c.execute(text("SELECT count(*) FROM business_items WHERE kind='card'")).scalar_one()==3
        assert c.execute(text("SELECT count(*) FROM business_events WHERE action='duplicate_acknowledged'")).scalar_one()==1


def test_policy_default_filter_status_export_and_pn_pagination(env,engine):
    with engine.begin() as c:
        catalog_fixture(c,env[1],env[3]['bm'],'A-PN',status='published')
        catalog_fixture(c,env[1],env[3]['bm'],'A-PN')
        for n in range(30):catalog_fixture(c,env[1],env[3]['bm'],f'PN-{n:03}')
    rule(engine,'nk.drafts.hide_published_duplicates',False)
    with client(env) as c:
        login(c)
        first=listed(c,env[1],sort='pn',status='local')
        second=listed(c,env[1],sort='pn',status='local',before_id=first['next_cursor'])
        rows=first['items']+second['items']
        assert len(rows)==31 and [r['body']['pn'] for r in rows]==sorted(r['body']['pn'] for r in rows)
        assert len(listed(c,env[1],status='published')['items'])==1
        exported=c.get('/api/catalog/export',params={'organization_id':str(env[1]),'status':'published'})
        assert len(list(load_workbook(BytesIO(exported.content)).active.values))==2


def test_schedule_single_job_interval_pause_and_delegate_revocation(env,engine):
    with client(env) as c:
        login(c,'administrator');connections(c,env,nk_key='fixture-key',nk_auto_sync=True)
    with ThreadPoolExecutor(max_workers=2) as pool:list(pool.map(lambda _:tick(engine,env[0]),range(2)))
    with engine.begin() as c:
        jobs=c.execute(text("SELECT * FROM business_items WHERE kind='sync'")).mappings().all()
        assert len(jobs)==1 and jobs[0]['body']['scheduled']
        assert jobs[0]['body']['policy']['values']['nk.sync.interval_seconds']==900
        c.execute(text("UPDATE business_items SET state='done' WHERE kind='sync'"))
    tick(engine,env[0])
    with engine.begin() as c:
        assert c.execute(text("SELECT count(*) FROM business_items WHERE kind='sync'")).scalar_one()==1
        c.execute(text("UPDATE business_items SET updated_at=now()-interval '1 day',state='blocked' WHERE kind='sync'"))
    tick(engine,env[0])
    with engine.begin() as c:
        assert c.execute(text("SELECT count(*) FROM business_items WHERE kind='sync'")).scalar_one()==1
        c.execute(text("UPDATE business_items SET state='done' WHERE kind='sync'"))
        c.execute(text('UPDATE users SET enabled=false WHERE id=:id'),{'id':env[3]['administrator']})
    tick(engine,env[0])
    with engine.connect() as c:
        assert c.execute(text("SELECT count(*) FROM business_items WHERE kind='sync'")).scalar_one()==1
        assert c.execute(text('SELECT count(*) FROM external_operations')).scalar_one()==0


def test_bulk_limit_partial_results_and_exact_retry(env,engine):
    with client(env) as c:
        login(c,'administrator');connections(c,env,writes_enabled=True)
    with client(env) as c:
        login(c,'specialist')
        ids=[c.post('/api/catalog',json=card(env,gtin='',pn=pn)).json()['id'] for pn in ['ABC','XYZ']]
        body={'organization_id':str(env[1]),'request_id':str(uuid4()),'action':'gtin','entries':[{'id':ids[0],'revision':1},{'id':ids[1],'revision':99}]}
        rule(engine,'nk.bulk.batch_size',1)
        assert c.post('/api/catalog-bulk',json=body).status_code==422
        rule(engine,'nk.bulk.batch_size',2)
        first=c.post('/api/catalog-bulk',json=body);assert first.status_code==200,first.text
        assert [r['ok'] for r in first.json()]==[True,False]
        rule(engine,'nk.bulk.batch_size',1)
        assert c.post('/api/catalog-bulk',json=body).json()==first.json()
        assert c.post('/api/catalog-bulk',json={**body,'action':'save'}).status_code==409
    with client(env) as c:
        login(c,'tsd');assert c.post('/api/catalog-bulk',json=body).status_code==403
    with engine.connect() as c:assert c.execute(text('SELECT count(*) FROM external_operations')).scalar_one()==1


def test_import_kind_preparation_preserves_exact_identity_and_source(env):
    book=Workbook();book.active.append(['PN','ТН ВЭД','Товарный знак']);book.active.append(['00123','8504408300','Actual Brand']);out=BytesIO();book.save(out);original=out.getvalue()
    with client(env) as c:
        login(c,'specialist');options=c.get('/api/catalog-options',params={'organization_id':str(env[1])}).json()
        kind=next(k for k in options['kinds'] if k['tnved']=='8504408300')
        r=c.post('/api/catalog-import/preview',params={'organization_id':str(env[1])},content=original).json()
        url='/api/catalog-import/'+r['id']
        prepared=c.post(url+'/prepare',json={'revision':1,'kinds':{'8504408300':kind['id']}})
        assert prepared.status_code==200,prepared.text
        d=prepared.json()['body']['rows'][0]['data']
        assert d['pn']=='00123' and d['brand']=='Actual Brand' and not d['gtin']
        assert d['category_id']==int(kind['category_id']) and d['attributes']['24003']=='КОЛИЧЕСТВО'
        assert c.get(url+'/source').content==original
        assert c.post(url+'/commit',json={'ids':[r['id']],'revision':2}).status_code==200


def test_pipeline_preflights_before_gtin_and_can_pause(env,engine,monkeypatch):
    with client(env) as c:
        login(c,'administrator');connections(c,env,writes_enabled=True,nk_key='fixture')
    with engine.begin() as c:
        actor=background(c,env[3]['specialist'],env[1]);a=catalog_fixture(c,env[1],actor.id,'A-PN')
        batch=store.create(actor,'import','Fixture',{'card_ids':[str(a['id'])]},env[1],state='done')
    monkeypatch.setattr(Client,'schema',lambda *a,**k:[{'attr_id':13914,'attr_type':'m','attr_name':'PN'},{'attr_id':99,'attr_type':'m','attr_name':'Required'}])
    with client(env) as c:
        login(c,'specialist');url='/api/catalog-import/'+str(batch['id'])
        assert c.post(url+'/pipeline',json={'revision':1,'agreement':True,'publish':True}).status_code==200
        worker=BusinessWorker(engine,env[0]);worker.pipeline_tick()
        with engine.connect() as db:
            assert db.execute(text('SELECT count(*) FROM external_operations')).scalar_one()==0
            saved=store.item(db,batch['id']);assert saved['state']=='blocked' and saved['body']['pipeline_counts']['errors']==1
        started=c.post(url+'/pipeline',json={'revision':saved['revision'],'agreement':True,'publish':True}).json()
        assert c.post(url+'/pause',json={'revision':started['revision']}).status_code==200
        worker.pipeline_tick()
        with engine.connect() as db:assert store.item(db,batch['id'])['state']=='blocked'


def test_publication_pipeline_waits_refreshes_then_queues_signature(env,engine,monkeypatch):
    with client(env) as c:
        login(c,'administrator');connections(c,env,writes_enabled=True,nk_key='fixture')
    with engine.begin() as c:
        actor=background(c,env[3]['specialist'],env[1]);a=catalog_fixture(c,env[1],actor.id,'PN-A',status='moderation')
        d=a['body'];d['external']['good_id']=42;d['external']['raw']['good_id']=42
        a=store.save(actor,a,'fixture',body=d)
        batch=store.create(actor,'import','Fixture',{'card_ids':[str(a['id'])],'pipeline_actor':str(actor.id),'pipeline':True,'publish':True,'results':{}},env[1],state='processing')
    raw={'good_id':42,'good_detailed_status':['notsigned'],'good_attrs':[{'attr_id':13914,'attr_value':'PN-A'}]}
    monkeypatch.setattr(Client,'details',lambda *a,**k:raw)
    worker=BusinessWorker(engine,env[0]);worker.pipeline_tick();worker.pipeline_tick()
    with engine.connect() as c:
        op=c.execute(text('SELECT kind,payload FROM external_operations')).mappings().one()
        assert op['kind']=='nk.document' and json.loads(bytes(op['payload']))['publicationAgreement'] is True
        assert store.item(c,batch['id'])['state']=='processing'
    worker.pipeline_tick()
    with engine.connect() as c:assert c.execute(text('SELECT count(*) FROM external_operations')).scalar_one()==1


def test_confirmed_gtin_is_available_in_common_product_registry(env,engine,monkeypatch):
    from marking_erp import operations
    from marking_erp.business.contracts import VERSION
    gtin='04640638349704'
    with client(env) as c:
        login(c,'administrator');connections(c,env,writes_enabled=True,nk_key='fixture')
    with client(env) as c:
        login(c,'specialist');created=c.post('/api/catalog',json=card(env,gtin='',pn='Exact-PN')).json()
        queued=c.post('/api/catalog/'+created['id']+'/gtin',json={'revision':1}).json()
    monkeypatch.setattr(Client,'call',lambda self,method,path,params=None,body=None:{'drafts':[{'gtin':gtin}]})
    worker=BusinessWorker(engine,env[0])
    assert operations.dispatch_once(engine,'fixture',{('nk.gtin',VERSION):worker.dispatch},allowed_kinds=['nk.gtin'])
    worker.reconcile(UUID(queued['operation_id']))
    with client(env) as c:
        login(c,'specialist');products=c.get('/api/products',params={'organization_id':str(env[1]),'q':'Exact-PN'})
        assert products.status_code==200,products.text
        assert products.json()['products'][0]['gtin']==gtin
    with engine.connect() as c:assert c.execute(text('SELECT count(*) FROM external_operations')).scalar_one()==1


def test_explicit_discard_preserves_local_history_and_restores_remote_identity(env,engine,monkeypatch):
    with client(env) as c:
        login(c,'administrator');connections(c,env,nk_key='fixture')
    with engine.begin() as c:
        row=catalog_fixture(c,env[1],env[3]['specialist'],'Local-PN',status='published',raw_pn='Remote-PN')
        actor=background(c,env[3]['specialist'],env[1]);data=row['body'];data['external']['good_id']=7
        row=store.save(actor,row,'fixture',body=data)
    monkeypatch.setattr(Client,'details',lambda *a,**k:{'good_id':7,'good_name':'Remote name','good_status':'published','good_attrs':[{'attr_id':13914,'attr_value':'Remote-PN'}]})
    with client(env) as c:
        login(c,'specialist');url='/api/catalog/'+str(row['id'])
        assert c.post(url+'/refresh',json={'revision':row['revision']}).status_code==409
        assert c.post(url+'/discard',json={'revision':row['revision']}).status_code==422
        discarded=c.post(url+'/discard',json={'revision':row['revision'],'agreement':True});assert discarded.status_code==200,discarded.text
        assert discarded.json()['body']['pn']=='Remote-PN' and discarded.json()['title']=='Remote-PN'
        events=c.get(url).json()['events']
        assert next(e for e in events if e['action']=='local_changes_discarded')['body']['previous']['pn']=='Local-PN'
