"""Spreadsheet output keeps literal strings and namespace compatibility metadata."""
from io import BytesIO
import zipfile
from openpyxl import Workbook,load_workbook
from marking_erp.business import catalog_excel


def test_formula_looking_strings_and_compatibility_namespaces_survive_safe_patch():
    book=Workbook();book.active.append(['PN','ТН ВЭД','Note']);book.active.append(['PN-1','8542399000','Literal'])
    book.active['C2']='=HYPERLINK("https://example.invalid","literal")';book.active['C2'].data_type='s'
    stream=BytesIO();book.save(stream)
    source=zipfile.ZipFile(BytesIO(stream.getvalue()));out=BytesIO()
    with zipfile.ZipFile(out,'w') as z:
        for info in source.infolist():
            data=source.read(info)
            if info.filename=='xl/worksheets/sheet1.xml':data=data.replace(b'<worksheet ',b'<worksheet xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" mc:Ignorable="x14ac" ',1)
            z.writestr(info,data)
    original=out.getvalue();meta=catalog_excel.read(original)
    result=catalog_excel.append_results(original,meta,{2:['04640638349704','=DDE-like text']})
    worksheet=load_workbook(BytesIO(result)).active
    assert worksheet['C2'].data_type==worksheet['E2'].data_type=='s'
    assert worksheet['E2'].value=='=DDE-like text'
    assert worksheet['D2'].value=='04640638349704'
    xml=zipfile.ZipFile(BytesIO(result)).read('xl/worksheets/sheet1.xml')
    assert b'xmlns:x14ac=' in xml and b'mc:Ignorable="x14ac"' in xml
