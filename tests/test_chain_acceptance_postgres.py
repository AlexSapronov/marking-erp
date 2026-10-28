"""API/queue/HTTP boundary acceptance. No business-state seeding or real integrations."""
import base64
from datetime import date
from io import BytesIO
from uuid import uuid4
import httpx
import pytest
from openpyxl import Workbook,load_workbook
from sqlalchemy import text
from test_auth_postgres import env,client,login
from test_signer_postgres import certificate
from test_documents_postgres import body as invoice
from chain_harness import Chain,SyntheticServices,PN,GTIN
from marking_erp.business.errors import RemoteError
from marking_erp.business.printing import register_printer

pytestmark=[pytest.mark.postgres,pytest.mark.usefixtures('fresh_schema')]


@pytest.fixture
def chain(env,engine,certificate,monkeypatch,tmp_path):
    monkeypatch.setattr('marking_erp.business.worker.BusinessWorker.start',lambda self:None)
    env[0].file_storage_path=tmp_path/'files'
    remote=SyntheticServices();transport=httpx.MockTransport(remote);real=httpx.Client
    monkeypatch.setattr(httpx,'Client',lambda *a,**kw:real(*a,transport=transport,**kw))
    value=Chain(env,engine,certificate,remote);value.configure();return value


def create_invoice(chain,c):
    value=invoice(chain.env,number='ACCEPTANCE-001',lines=[{'pn':PN,'gtin':GTIN,'tnved':'8542399000','quantity':10,'labels':1}])
    doc=chain.response(c,'/api/documents',value)
    chain.response(c,'/api/documents/'+doc['id']+'/prepare',{'revision':1})
    body={'source_id':doc['id'],'revision':2,'request_id':str(uuid4()),'confirmed':True}
    item=chain.response(c,'/api/marking/start',body)
    assert chain.response(c,'/api/marking/start',body)['id']==item['id']
    return doc['id'],item['id']


def full_excel():
    book=Workbook();book.active.append(['PN','GTIN','Количество'])
    book.active.append([PN,GTIN,10]);book.active.append(['UNMARKED-001','',7])
    out=BytesIO();book.save(out);return out.getvalue()


def test_entire_chain_from_nk_to_archived_receipt_and_shipment(chain,engine):
    env=chain.env
    with client(env) as bm,client(env) as logistics,client(env) as specialist,client(env) as warehouse:
        login(bm);login(logistics,'logistics');login(specialist,'specialist');login(warehouse,'warehouse')
        catalog=chain.publish(bm);source,item=create_invoice(chain,bm);chain.pump()
        row=chain.row(bm,item);assert row['state']=='labels_ready'
        data=chain.response(bm,'/api/business/'+item,method='GET');code=data['codes'][0]
        with engine.connect() as c:raw=bytes(c.execute(text('SELECT payload FROM business_codes WHERE id=:id'),{'id':code['id']}).scalar_one()).decode()
        kit=chain.response(bm,'/api/business-files/'+row['body']['label_file_id'],method='GET');assert kit[:2]==b'PK'
        # One exact-kit intent, followed by a separate authenticated printer agent.
        with engine.begin() as c:device,token=register_printer(c,env[1],'Acceptance printer','Synthetic driver boundary')
        chain.response(specialist,'/api/business/'+item+'/auto-print',{'request_id':str(uuid4()),'printer_id':str(device),'revision':row['revision'],'confirmed':True})
        from marking_erp.business.auto_print import tick
        tick(engine)
        with client(env) as agent:
            headers={'Authorization':'Bearer '+token}
            job=chain.response(agent,'/api/print-agent/claim',headers=headers)['job']
            from hashlib import sha256
            pdf=base64.b64decode(job['pdf_base64']);assert pdf.startswith(b'%PDF') and sha256(pdf).hexdigest()==job['pdf_sha256']
            receipt={'job_id':job['id'],'lease':job['lease'],'pdf_sha256':job['pdf_sha256'],'outcome':'spooled','spool_reference':'synthetic-driver-only'}
            chain.response(agent,'/api/print-agent/result',receipt,headers=headers)
            chain.response(agent,'/api/print-agent/result',receipt,headers=headers)
            assert chain.response(agent,'/api/print-agent/claim',headers=headers)['job'] is None
        # Exact-PDF reprint, copy limit, post-print scan, lost scan reply replay.
        rid=str(uuid4());url='/api/business-codes/'+code['id']+'/reprint'
        chain.response(specialist,url,{'request_id':rid,'reason':'Synthetic label readability check','copies':2},status=422)
        chain.response(specialist,url,{'request_id':rid,'reason':'Synthetic label readability check','copies':1})
        chain.response(specialist,'/api/reprints/'+rid+'/printed')
        control={'request_id':str(uuid4()),'raw':raw}
        assert chain.response(specialist,'/api/reprints/'+rid+'/scans',control)['state']=='completed'
        assert chain.response(specialist,'/api/reprints/'+rid+'/scans',control)['scanned_copies']==1
        # Application is blocked until handoff, and a report rejection is visible.
        chain.response(bm,'/api/business/'+item+'/applied',{'revision':row['revision'],'confirmed':True},status=409)
        chain.response(bm,'/api/business/'+item+'/labels-handoff',{'request_id':str(uuid4()),'revision':row['revision'],'file_id':row['body']['label_file_id'],'confirmed':True})
        chain.action(bm,item,'applied',confirmed=True)
        chain.action(bm,item,'report',confirmed=True,production_date=date.today().isoformat());chain.pump()
        assert chain.row(bm,item)['state']=='rejected'
        chain.action(bm,item,'retry-report',confirmed=True,production_date=date.today().isoformat());chain.pump()
        assert chain.row(bm,item)['state']=='applied'
        # ATK registration uses real worker dispatch and independent registry readback.
        ids=chain.response(bm,'/api/atks',{'organization_id':str(env[1]),'ids':[item],'request_id':str(uuid4())})['ids'];assert len(ids)==1
        atk=ids[0];chain.pump();assert chain.row(bm,atk)['state']=='ready'
        chain.action(logistics,atk,'transferred',confirmed=True)
        day=date.today();number='12345678/'+day.strftime('%d%m%y')+'/1234567'
        fts=chain.response(logistics,'/api/fts',{'organization_id':str(env[1]),'request_id':str(uuid4()),'atks':[atk],'declaration_number':number,'declaration_date':day.isoformat(),'country_of_origin':'China','conformity_document':'Synthetic acceptance permit'})['id']
        chain.response(logistics,'/api/fts/'+fts+'/submit',{'revision':1,'confirmed':True})
        chain.remote.hold_introduced=True
        with pytest.raises(RemoteError):chain.pump()
        assert chain.row(logistics,fts)['state']=='fts_submitted'
        chain.remote.introduce();chain.pump();assert chain.row(logistics,fts)['state']=='circulation_confirmed'
        # Logistics returns the same invoice to BM; marked codes remain unchanged.
        row=chain.row(bm,item)
        chain.response(logistics,'/api/logistics/documents/'+item+'/request-composition',{'revision':row['revision'],'confirmation':'ACCEPTANCE-001'})
        doc=chain.response(bm,'/api/documents/'+source,method='GET')['document']
        params={'revision':doc['marking_revision'],'source_revision':doc['revision']};content=full_excel()
        preview=chain.response(bm,'/api/documents/'+source+'/warehouse-composition/preview',params=params,content=content)
        assert not preview['errors'],preview
        params['sha256']=preview['sha256']
        chain.response(bm,'/api/documents/'+source+'/warehouse-composition/commit',params=params,content=content)
        assert chain.response(bm,'/api/documents/'+source+'/warehouse-composition/commit',params=params,content=content)['replayed']
        doc=chain.response(bm,'/api/documents/'+source,method='GET')['document'];assert len(doc['lines'])==2
        tasks=[]
        for kind in ('receipt','shipment'):
            if kind=='shipment':
                # Damage discovered after receipt: emit, apply/report, then submit
                # LK_REMARK. Only official readback retires the original code.
                remark=chain.response(specialist,'/api/remarks',{'organization_id':str(env[1]),'request_id':str(uuid4()),'old_code_id':code['id'],'reason':'Damaged label after warehouse receipt','confirmed':True})['id']
                chain.pump();assert chain.row(specialist,remark)['state']=='labels_ready'
                for action,extra in [('applied',{}),('report',{'production_date':day.isoformat()}),('submit',{'remarking_date':day.isoformat()})]:
                    latest=chain.row(specialist,remark)
                    chain.response(specialist,'/api/remarks/'+remark+'/'+action,{'revision':latest['revision'],'confirmed':True,**extra})
                    chain.pump()
                assert chain.row(specialist,remark)['state']=='done'
                chain.response(specialist,url,{'request_id':str(uuid4()),'reason':'Old marking is now retired'},status=409)
                with engine.connect() as c:
                    assert c.execute(text('SELECT state FROM business_codes WHERE id=:id'),{'id':code['id']}).scalar_one()=='withdrawn'
                    raw=bytes(c.execute(text('SELECT payload FROM business_codes WHERE item_id=:id'),{'id':remark}).scalar_one()).decode()
            tid=str(uuid4());body={'request_id':tid,'source_id':source,'source_revision':doc['revision'],'confirmed':True,'kind':kind,'assignees':[str(env[3]['warehouse'])]}
            chain.response(logistics,'/api/logistics/warehouse-transfer',body)
            assert chain.response(logistics,'/api/logistics/warehouse-transfer',body)['replayed']
            scan={'request_id':str(uuid4()),'raw':raw}
            chain.response(warehouse,'/api/warehouse/tasks/'+tid+'/scans',scan)
            chain.response(warehouse,'/api/warehouse/tasks/'+tid+'/scans',{'request_id':str(uuid4()),'raw':raw},status=409)
            current=chain.response(warehouse,'/api/warehouse/tasks/'+tid,method='GET')['task']
            if kind=='receipt':
                cis=raw.split('\x1d')[0][:38];chain.remote.codes[cis]['ownerInn']='7800000000'
                denied=chain.response(warehouse,'/api/warehouse/tasks/'+tid+'/check',{'request_id':str(uuid4()),'revision':current['revision']})
                assert not denied['ok'] and not denied['external_verified']
                current=chain.response(warehouse,'/api/warehouse/tasks/'+tid,method='GET')['task']
                chain.response(warehouse,'/api/warehouse/tasks/'+tid+'/close',{'revision':current['revision'],'confirmation':'ACCEPTANCE-001'},status=409)
                chain.remote.codes[cis]['ownerInn']='7700000000'
            check=chain.response(warehouse,'/api/warehouse/tasks/'+tid+'/check',{'request_id':str(uuid4()),'revision':current['revision']})
            assert check['ok'] and check['external_verified'],check
            current=chain.response(warehouse,'/api/warehouse/tasks/'+tid,method='GET')['task']
            chain.response(warehouse,'/api/warehouse/tasks/'+tid+'/close',{'revision':current['revision'],'confirmation':'ACCEPTANCE-001'})
            assert chain.response(warehouse,'/api/warehouse/tasks/'+tid+'/scans',scan)['replayed']
            tasks.append(tid)
        exported=chain.response(warehouse,'/api/warehouse/tasks/'+tasks[1]+'/export',method='GET')
        book=load_workbook(BytesIO(exported));assert list(book['Немаркированные строки'].values)[1][1]==7
        ready=chain.response(logistics,'/api/logistics/documents/'+item,method='GET');assert ready['can_complete'],ready
        finish={'revision':ready['document']['revision'],'confirmation':'ACCEPTANCE-001'}
        chain.response(logistics,'/api/logistics/documents/'+item+'/complete',finish)
        assert chain.response(logistics,'/api/logistics/documents/'+item+'/complete',finish)['replayed']
        with client(env) as admin:
            login(admin,'administrator')
            for tid in tasks:
                current=chain.response(admin,'/api/warehouse/tasks/'+tid,method='GET')['task']
                chain.response(admin,'/api/warehouse/tasks/'+tid+'/archive',{'revision':current['revision'],'confirmation':'ACCEPTANCE-001'})
        with engine.connect() as c:
            assert c.execute(text('SELECT count(*) FROM warehouse_tasks WHERE archived_at IS NOT NULL')).scalar_one()==2
            assert c.execute(text("SELECT count(*) FROM external_operations WHERE state NOT IN ('confirmed','rejected','cancelled_before_dispatch')")).scalar_one()==0
            assert c.execute(text('SELECT count(*) FROM business_links WHERE NOT applied')).scalar_one()==0 # Every receipt was consumed, including the rejection.
            assert c.execute(text('SELECT count(*) FROM business_codes')).scalar_one()==2
            assert c.execute(text('SELECT count(*) FROM code_replacements')).scalar_one()==1
            assert c.execute(text('SELECT count(*) FROM print_jobs')).scalar_one()==1
            assert c.execute(text('SELECT count(DISTINCT (product_id,pn,tnved)) FROM product_pns WHERE pn=:pn'),{'pn':PN}).scalar_one()==1
        flow=chain.response(bm,'/api/workflow/document/'+source,method='GET');assert flow['stage']=='Завершено логистом'
        assert any(link['target']['id']==catalog for group in flow['groups'] for link in group['items'])


def test_lost_order_response_is_unknown_and_never_reemitted(chain,engine):
    env=chain.env
    with client(env) as bm:
        login(bm);chain.publish(bm);source,item=create_invoice(chain,bm)
        chain.remote.lose_order_reply=True;chain.pump()
        with engine.connect() as c:
            op=c.execute(text("SELECT * FROM external_operations WHERE kind='bm.order'")).mappings().one()
            assert op['state']=='unknown'
            attempts=c.execute(text('SELECT count(*) FROM operation_attempts WHERE operation_id=:id'),{'id':op['id']}).scalar_one()
        chain.response(bm,'/api/business-operations/'+str(op['id'])+'/resume',{'revision':chain.row(bm,item)['revision'],'confirmed':True},status=409)
        count=chain.remote.calls.count(('POST','/api/v3/order'))
        chain.pump();chain.pump()
        assert chain.remote.calls.count(('POST','/api/v3/order'))==count==1
        flow=chain.response(bm,'/api/workflow/document/'+source,method='GET')
        assert 'неизвестный результат' in flow['next_action']
        with engine.connect() as c:
            assert c.execute(text('SELECT count(*) FROM operation_attempts WHERE operation_id=:id'),{'id':op['id']}).scalar_one()==attempts==1
            assert c.execute(text('SELECT count(*) FROM business_codes')).scalar_one()==0
