"""Local workflow evidence fixtures; never performs real signing or emission."""
from io import BytesIO
from uuid import uuid4
import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from fastapi import HTTPException
from test_auth_postgres import env,client,login
from test_business_registry_postgres import allocated
from test_documents_postgres import body as invoice
from marking_erp.business import store
from marking_erp.business.worker import background
from marking_erp.business.supplier_batches import require_ready
from marking_erp.business.marking import atk_export,apply

pytestmark=[pytest.mark.postgres,pytest.mark.usefixtures('fresh_schema')]


def evidence(env,engine,item,code,*,fts=True):
    with engine.begin() as c:
        actor=background(c,env[3]['specialist'],env[1])
        c.execute(text("UPDATE business_items SET state='applied' WHERE id=:id"),{'id':item})
        c.execute(text("UPDATE business_codes SET state='applied' WHERE item_id=:id"),{'id':item})
        atk=store.create(actor,'atk','ATK fixture',{'supplier':'Поставщик А','tnved':'8542399000','atk':'1'*25,'documents':[str(item)],'codes':[str(code['id'])]},env[1],state='transferred')
        if fts:store.create(actor,'fts','FTS fixture',{'atks':[str(atk['id'])],'external_document_id':str(uuid4()),'codes':[{'id':str(code['id'])}]},env[1],state='circulation_confirmed')
        return atk['id']


def receipt(env,source):
    with client(env) as c:
        login(c,'logistics');id=uuid4()
        r=c.post('/api/logistics/warehouse-transfer',json={'request_id':str(id),'source_id':str(source),'source_revision':2,'kind':'receipt','assignees':[str(env[3]['warehouse'])]});assert r.status_code==200,r.text
        return id


def test_completion_requires_real_evidence_and_permission(allocated,engine):
    env,item,source,code=allocated
    with client(env) as c:
        login(c,'bm')
        assert c.get('/api/logistics/documents/'+str(item)).status_code==403
        assert c.post('/api/logistics/documents/'+str(item)+'/complete',json={'revision':2,'confirmation':'INV-01'}).status_code==403
    with client(env) as c:
        login(c,'logistics');data=c.get('/api/logistics/documents/'+str(item)).json()
        assert not data['can_complete'] and len(data['blockers'])>=4
        response=c.post('/api/logistics/documents/'+str(item)+'/complete',json={'revision':2,'confirmation':'INV-01'})
        assert response.status_code==409
    evidence(env,engine,item,code,fts=False);task=receipt(env,source)
    with engine.begin() as c:c.execute(text("UPDATE warehouse_tasks SET state='closed' WHERE id=:id"),{'id':task})
    with client(env) as c:
        login(c,'logistics');data=c.get('/api/logistics/documents/'+str(item)).json()
        assert any('ФТС' in x for x in data['blockers'])
        assert any('физическую' in x for x in data['blockers'])


def test_complete_replay_freezes_workflow_then_privileged_reopen(allocated,engine):
    env,item,source,code=allocated;evidence(env,engine,item,code);task=receipt(env,source)
    with engine.begin() as c:c.execute(text("UPDATE warehouse_tasks SET state='closed',verified_at=now() WHERE id=:id"),{'id':task})
    with client(env) as c:
        login(c,'logistics');data=c.get('/api/logistics/documents/'+str(item)).json();assert data['can_complete'],data
        body={'revision':data['document']['revision'],'confirmation':'INV-01'}
        assert c.post('/api/logistics/documents/'+str(item)+'/complete',json=body).status_code==200
        assert c.post('/api/logistics/documents/'+str(item)+'/complete',json=body).json()['replayed']
        listing=c.get('/api/logistics/documents',params={'organization_id':str(env[1]),'completed':True}).json();assert listing['documents'][0]['id']==str(item)
        latest=c.get('/api/logistics/documents/'+str(item)).json()['document']
        assert c.post('/api/logistics/documents/'+str(item)+'/reopen',json={'revision':latest['revision'],'confirmation':'INV-01','reason':'Need document correction'}).status_code==403
    with client(env) as c:
        login(c,'specialist')
        assert c.post('/api/warehouse/tasks/'+str(task)+'/reopen',json={'revision':1,'confirmation':'INV-01','reason':'Need document correction'}).status_code==409
        assert c.post('/api/documents/'+str(source)+'/archive',json={'revision':2,'confirmation':'INV-01','reason':'Need document correction'}).status_code==409
        assert c.post('/api/logistics/documents/'+str(item)+'/reopen',json={'revision':latest['revision'],'confirmation':'INV-01','reason':'Need document correction'}).status_code==200
    with engine.connect() as c:
        assert c.execute(text("SELECT count(*) FROM business_events WHERE item_id=:id AND action='logistics_completed'"),{'id':item}).scalar_one()==1
        assert c.execute(text("SELECT count(*) FROM business_events WHERE item_id=:id AND action='logistics_reopened'"),{'id':item}).scalar_one()==1


def test_supplier_batch_waits_all_invoices_and_not_first_page(allocated,engine):
    env,item,source,code=allocated
    with engine.begin() as c:c.execute(text("UPDATE business_items SET state='applied' WHERE id=:id"),{'id':item})
    with client(env) as c:
        login(c)
        for n in range(26):
            r=c.post('/api/documents',json=invoice(env,number='INV-'+str(n+2)));assert r.status_code==200,r.text
        batch=c.get('/api/atks/supplier-batch',params={'organization_id':str(env[1]),'supplier':'Поставщик А'}).json()
        assert not batch['ready'] and len(batch['documents'])==27
        result=c.post('/api/atks',json={'organization_id':str(env[1]),'ids':[str(item)],'request_id':str(uuid4())})
        assert result.status_code==409 and 'всем' in result.json()['detail']
    with engine.begin() as c:
        with pytest.raises(HTTPException):require_ready(c,env[1],'Поставщик А',{item})


def test_atk_excel_contains_only_its_bucket_and_transformation_refreshes_export(allocated,engine):
    env,item,source,code=allocated
    with engine.begin() as c:
        actor=background(c,env[3]['specialist'],env[1])
        other=uuid4()
        # A different bucket in the same invoice must never leak into this XLSX.
        c.execute(text('''INSERT INTO business_codes(id,organization_id,cis,gtin,base,payload,item_id,line_index,quantity,state,tnved,pn,evidence_ref)
            SELECT :id,organization_id,substring(cis,1,37)||'X',gtin,base,payload,item_id,1,quantity,state,'8542399001','PN-OTHER','fixture' FROM business_codes WHERE id=:source'''),{'id':other,'source':code['id']})
        row=store.create(actor,'atk','ATK fixture',{'supplier':'Поставщик А','tnved':'8542399000','atk':'1'*25,'documents':[str(item)],'codes':[str(code['id'])]},env[1],state='ready')
        first=atk_export(actor,row)
        saved=c.execute(text('SELECT body FROM business_artifacts WHERE id=:id'),{'id':first['body']['file_id']}).scalar_one()
        values=list(load_workbook(BytesIO(bytes(saved))).active.values)
        assert len(values)==2 and values[1][0]=='PN/A-1' and values[1][4]=='1'*25
        transform=store.create(actor,'transform','Remove fixture',{'atk_item':str(row['id']),'kind':'REMOVING','codes':[str(code['id'])]},env[1])
        apply(actor,transform,{'kind':'bm.transform'},{},{},{},None,env[0])
        exports=c.execute(text("SELECT version,body FROM business_artifacts WHERE item_id=:id AND kind='atk' ORDER BY version"),{'id':row['id']}).mappings().all()
        assert len(exports)==2 and exports[1]['version']==2
        assert len(list(load_workbook(BytesIO(bytes(exports[1]['body']))).active.values))==1
