"""Exact remote PN identity, visibility and pagination; no NK calls."""
from io import BytesIO
import pytest
from sqlalchemy import text
from openpyxl import load_workbook
from marking_erp import auth
from marking_erp.business import store
from test_auth_postgres import env,client,login

pytestmark=[pytest.mark.postgres,pytest.mark.usefixtures('fresh_schema')]


def catalog_fixture(c,org,owner,pn,*,status=None,raw_pn=None,state=None,title=None,extra_attrs=None):
    user=dict(c.execute(text('SELECT * FROM users WHERE id=:id'),{'id':owner}).mappings().one())
    current=auth.Actor(c,user,'fixture',None)
    data={'pn':pn,'name':'Component','gtin':'','tnved':'8542399000','brand':'Brand','attributes':{},'category_id':1,'dirty':status is None}
    if status:
        raw={'good_id':pn,'good_status':status,'good_detailed_status':[status],'brand_name':'Brand',
             'good_attrs':[{'attr_id':13914,'attr_value':raw_pn if raw_pn is not None else pn}]+(extra_attrs or [])}
        data['external']={'statuses':[status],'raw':raw}
        if raw_pn is not None:data['dirty']=True
    return store.create(current,'card',title or pn,data,org,state=state or ('active' if status else 'draft'))


def listed(c,org,**params):
    response=c.get('/api/catalog',params={'organization_id':str(org),**params})
    assert response.status_code==200,response.text
    return response.json()


def test_hide_include_only_export_use_exact_pn_and_do_not_mutate(env,engine):
    with engine.begin() as db:
        pub=catalog_fixture(db,env[1],env[3]['bm'],'MPQ5-S/1220A',status='published')
        draft=catalog_fixture(db,env[1],env[3]['bm'],'MPQ5-S/1220A')
        other=catalog_fixture(db,env[1],env[3]['bm'],'MPQ5-S/1220B')
        punctuation=catalog_fixture(db,env[1],env[3]['bm'],'MPQ5S/1220A')
        lowercase=catalog_fixture(db,env[1],env[3]['bm'],'mpq5-S/1220A')
        count=db.execute(text('SELECT count(*) FROM business_events')).scalar_one()
    with client(env) as c:
        login(c)
        assert {r['id'] for r in listed(c,env[1])['items']}=={str(r['id']) for r in (pub,other,punctuation,lowercase)}
        items=listed(c,env[1],draft_duplicates='include')['items']
        assert len(items)==5 and next(r for r in items if r['id']==str(draft['id']))['duplicate_draft'] is True
        assert listed(c,env[1],draft_duplicates='only')['items'][0]['id']==str(draft['id'])
        for mode,expected in [('hide',4),('include',5),('only',1)]:
            result=c.get('/api/catalog/export',params={'organization_id':str(env[1]),'draft_duplicates':mode})
            assert result.status_code==200,result.text
            rows=list(load_workbook(BytesIO(result.content)).active.values)
            assert len(rows)==expected+1
            if mode=='only':assert rows[1][-1]=='Есть опубликованная карточка с точным PN'
        assert c.get('/api/catalog/'+str(draft['id'])).status_code==200
        assert c.get('/api/catalog',params={'organization_id':str(env[1]),'draft_duplicates':'unsafe'}).status_code==422
    with engine.connect() as db:
        assert db.execute(text('SELECT count(*) FROM business_events')).scalar_one()==count
        assert db.execute(text('SELECT count(*) FROM external_operations')).scalar_one()==0
        assert db.execute(text('SELECT state FROM business_items WHERE id=:id'),{'id':draft['id']}).scalar_one()=='draft'


def test_remote_publication_pn_survives_local_edits_and_ambiguous_source_is_not_guessed(env,engine):
    with engine.begin() as db:
        edited=catalog_fixture(db,env[1],env[3]['bm'],'LOCAL-CHANGE',status='published',raw_pn='REMOTE-ORIGINAL')
        original=catalog_fixture(db,env[1],env[3]['bm'],'REMOTE-ORIGINAL')
        local=catalog_fixture(db,env[1],env[3]['bm'],'LOCAL-CHANGE')
        catalog_fixture(db,env[1],env[3]['bm'],'AMBIGUOUS',status='published',extra_attrs=[{'attr_id':'13914','attr_value':'OTHER-PN'}])
        ambiguous=catalog_fixture(db,env[1],env[3]['bm'],'AMBIGUOUS')
        # Merely being active or awaiting moderation is not a draft eligible for hiding.
        active=catalog_fixture(db,env[1],env[3]['bm'],'REMOTE-ORIGINAL',state='active')
        moderation=catalog_fixture(db,env[1],env[3]['bm'],'REMOTE-ORIGINAL',status='moderation')
    with client(env) as c:
        login(c)
        ids={r['id'] for r in listed(c,env[1])['items']}
        assert str(original['id']) not in ids
        assert {str(r['id']) for r in (edited,local,ambiguous,active,moderation)}.issubset(ids)


def test_hidden_drafts_are_filtered_before_keyset_pagination_and_pn_search(env,engine):
    with engine.begin() as db:
        pub=catalog_fixture(db,env[1],env[3]['bm'],'VISIBLE-PN',status='published',title='Published component')
        visible=[pub]
        for n in range(29):
            visible.append(catalog_fixture(db,env[1],env[3]['bm'],f'OTHER-{n:03}',title=f'Component {n}'))
            catalog_fixture(db,env[1],env[3]['bm'],'VISIBLE-PN')
    with client(env) as c:
        login(c)
        first=listed(c,env[1]);assert len(first['items'])==25 and first['next_cursor']
        second=listed(c,env[1],before_id=first['next_cursor']);assert len(second['items'])==5 and second['next_cursor'] is None
        assert {r['id'] for r in first['items']+second['items']}=={str(r['id']) for r in visible}
        assert [r['id'] for r in listed(c,env[1],q='VISIBLE-PN')['items']]==[str(pub['id'])]


def test_other_org_and_inaccessible_publication_do_not_hide_or_leak(env,engine):
    with engine.begin() as db:
        catalog_fixture(db,env[2],env[3]['bm'],'ISOLATED',status='published')
        own=catalog_fixture(db,env[1],env[3]['bm'],'ISOLATED')
        hidden_pub=catalog_fixture(db,env[1],env[3]['specialist'],'PRIVATE-PUB',status='published')
        visible_draft=catalog_fixture(db,env[1],env[3]['bm'],'PRIVATE-PUB')
        db.execute(text("UPDATE group_grants SET scope='own' WHERE group_slug='bm' AND permission_id='nk.cards.view'"))
    with client(env) as c:
        login(c)
        assert {r['id'] for r in listed(c,env[1])['items']}=={str(own['id']),str(visible_draft['id'])}
        assert not listed(c,env[1],draft_duplicates='only')['items']
        assert c.get('/api/catalog',params={'organization_id':str(env[2])}).status_code==403
        with engine.begin() as db:
            db.execute(text("UPDATE group_grants SET scope='assigned' WHERE group_slug='bm' AND permission_id='nk.cards.view'"))
            db.execute(text('INSERT INTO operation_assignments VALUES (:target,:user)'),{'target':visible_draft['id'],'user':env[3]['bm']})
        assert [r['id'] for r in listed(c,env[1])['items']]==[str(visible_draft['id'])]
        with engine.begin() as db:db.execute(text('INSERT INTO operation_assignments VALUES (:target,:user)'),{'target':hidden_pub['id'],'user':env[3]['bm']})
        assert [r['id'] for r in listed(c,env[1],draft_duplicates='only')['items']]==[str(visible_draft['id'])]


def test_archive_is_complete_and_does_not_change_remote_publication_truth(env,engine):
    with engine.begin() as db:
        archived_pub=catalog_fixture(db,env[1],env[3]['bm'],'ARCHIVED-PUB',status='published',state='archived')
        catalog_fixture(db,env[1],env[3]['bm'],'ARCHIVED-PUB')
        archived_draft=catalog_fixture(db,env[1],env[3]['bm'],'ARCHIVED-PUB',state='archived')
    with client(env) as c:
        login(c)
        assert not listed(c,env[1])['items']
        items=listed(c,env[1],archived=True,draft_duplicates='only')['items']
        assert {r['id'] for r in items}=={str(archived_pub['id']),str(archived_draft['id'])}
        result=c.get('/api/catalog/export',params={'organization_id':str(env[1]),'archived':True})
        assert len(list(load_workbook(BytesIO(result.content)).active.values))==3


def test_remote_drafts_and_alternate_model_attribute_use_confirmed_identity(env,engine):
    import json
    with engine.begin() as db:
        pub=catalog_fixture(db,env[1],env[3]['bm'],'FALLBACK-1',status='published')
        raw=pub['body']['external']['raw']
        raw['good_attrs']=[{'attr_id':900,'attr_name':'Модель / артикул','attr_value':'FALLBACK-1'}]
        body={**pub['body'],'external':{'statuses':[],'raw':raw}}
        db.execute(text('UPDATE business_items SET body=CAST(:body AS jsonb) WHERE id=:id'),{'id':pub['id'],'body':json.dumps(body)})
        draft=catalog_fixture(db,env[1],env[3]['bm'],'FALLBACK-1',status='draft')
        correction=catalog_fixture(db,env[1],env[3]['bm'],'FALLBACK-1',status='errors')
        unsigned=catalog_fixture(db,env[1],env[3]['bm'],'FALLBACK-1',status='notsigned')
        # The official PN attribute wins over an unrelated model field.
        official=catalog_fixture(db,env[1],env[3]['bm'],'OFFICIAL-PN',status='published',extra_attrs=[{'attr_id':900,'attr_name':'PN','attr_value':'NOT-THE-PN'}])
        catalog_fixture(db,env[1],env[3]['bm'],'OFFICIAL-PN')
        other=catalog_fixture(db,env[1],env[3]['bm'],'NOT-THE-PN')
    with client(env) as c:
        login(c)
        ids={r['id'] for r in listed(c,env[1])['items']}
        assert ids=={str(r['id']) for r in (pub,unsigned,official,other)}
        hidden={r['id'] for r in listed(c,env[1],draft_duplicates='only')['items']}
        assert str(draft['id']) in hidden and str(correction['id']) in hidden
