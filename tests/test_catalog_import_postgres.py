"""Source workbook, exact row links, retries and export permissions through the API."""
from io import BytesIO
import json,zipfile
from uuid import uuid4
import pytest
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
from openpyxl.comments import Comment
from sqlalchemy import text
from test_auth_postgres import env,client,login
from test_business_postgres import card

pytestmark=[pytest.mark.postgres,pytest.mark.usefixtures('fresh_schema')]
GTIN='04640638349704'


def source_book():
    book=Workbook();sheet=book.active;sheet.title='Товары'
    sheet.append(['PN','ТН ВЭД','Товарный знак','ID категории','Примечание'])
    sheet.append(['00123','8542399000','Brand',1,'Keep original'])
    sheet.append([None]*5)
    sheet.append(['00123','8542399000','Brand',1,'Repeated source row'])
    sheet.append(['00123','8542399000','Other Brand',1,'Different manufacturer'])
    sheet['A2'].font=Font(bold=True,color='FF123456');sheet['B2'].fill=PatternFill('solid',fgColor='FFABCDEF')
    sheet['A2'].comment=Comment('Original note','Warehouse');sheet.freeze_panes='B2';sheet.column_dimensions['A'].width=24
    note=book.create_sheet('Справка');note['A1']='Unchanged second sheet';note.merge_cells('A1:C1')
    out=BytesIO();book.save(out);return out.getvalue()


def preview(c,env,content=None,**params):
    result=c.post('/api/catalog-import/preview',params={'organization_id':str(env[1]),'filename':'Source.xlsx',**params},content=content or source_book())
    assert result.status_code==200,result.text
    return result.json()


def commit(c,created,reason='Confirmed different manufacturer'):
    body={'ids':[created['id']],'revision':created['revision'],'reason':reason}
    result=c.post('/api/catalog-import/'+created['id']+'/commit',json=body)
    assert result.status_code==200,result.text
    return result.json(),body


def test_workbook_rows_styles_notes_and_other_members_survive_with_one_card_per_exact_identity(env,engine):
    original=source_book()
    with client(env) as c:
        login(c);created=preview(c,env,original)
        assert [r['row'] for r in created['rows']]==[2,4,5]
        assert created['rows'][1]['duplicate_row']==2
        result,body=commit(c,created)
        assert len(result['card_ids'])==2 and len(result['row_links'])==3
        assert result['row_links'][0]['card_id']==result['row_links'][1]['card_id']
        assert result['row_links'][2]['card_id']!=result['row_links'][0]['card_id']
        assert c.post('/api/catalog-import/'+created['id']+'/commit',json=body).json()==result
        assert c.get('/api/catalog-import/'+created['id']+'/source').content==original
        with engine.begin() as db:
            # Local fixture values, no request to NK or fake external completion.
            db.execute(text("UPDATE business_items SET body=jsonb_set(body,'{gtin}',to_jsonb(CAST(:gtin AS text))) WHERE id=:id"),{'id':result['card_ids'][0],'gtin':GTIN})
        response=c.get('/api/catalog-import/'+created['id']+'/export');assert response.status_code==200,response.text
        output=response.content;book=load_workbook(BytesIO(output));sheet=book['Товары']
        assert book.sheetnames==['Товары','Справка']
        assert sheet['A2'].value=='00123' and sheet['A4'].value=='00123' and sheet['A5'].value=='00123'
        assert sheet['F2'].value==sheet['F4'].value==GTIN and sheet['F2'].data_type=='s'
        assert sheet['F5'].value in ('',None) and sheet['F3'].value is None
        assert sheet['E2'].value=='Keep original' and sheet['E4'].value=='Repeated source row'
        assert sheet['A2'].font.bold and sheet['A2'].font.color.rgb=='FF123456'
        assert sheet['B2'].fill.fgColor.rgb=='FFABCDEF' and sheet['A2'].comment.text=='Original note'
        assert sheet.freeze_panes=='B2' and sheet.column_dimensions['A'].width==24
        before=zipfile.ZipFile(BytesIO(original));after=zipfile.ZipFile(BytesIO(output))
        changed=[name for name in before.namelist() if before.read(name)!=after.read(name)]
        assert changed==['xl/worksheets/sheet1.xml']
        assert before.namelist()==after.namelist()
    with engine.connect() as db:
        assert db.execute(text("SELECT count(*) FROM business_items WHERE kind='card'")).scalar_one()==2
        assert db.execute(text('SELECT count(*) FROM external_operations')).scalar_one()==0


def test_numeric_pn_and_invalid_category_are_row_errors_and_commit_is_atomic(env,engine):
    book=Workbook();sheet=book.active;sheet.append(['PN','ТН ВЭД','ID категории'])
    sheet.append(['VALID','8542399000',1]);sheet.append([123,'8542399000',1]);sheet.append(['NEXT','8542399000','bad'])
    stream=BytesIO();book.save(stream)
    with client(env) as c:
        login(c);created=preview(c,env,stream.getvalue())
        assert 'текстом' in created['rows'][1]['error'] and created['rows'][2]['error']
        response=c.post('/api/catalog-import/'+created['id']+'/commit',json={'ids':[created['id']],'revision':1})
        assert response.status_code==422
    with engine.connect() as db:assert db.execute(text("SELECT count(*) FROM business_items WHERE kind='card'")).scalar_one()==0


def test_source_and_result_require_export_scope_including_generic_file_route(env,engine):
    with client(env) as specialist,client(env) as bm,client(env) as tsd:
        login(specialist,'specialist');created=preview(specialist,env);commit(specialist,created)
        login(bm);login(tsd,'tsd')
        with engine.begin() as db:
            file=db.execute(text("SELECT id FROM business_artifacts WHERE item_id=:id AND kind='catalog_import_source'"),{'id':created['id']}).scalar_one()
            db.execute(text("UPDATE group_grants SET scope='all' WHERE group_slug='bm' AND permission_id='split.operations.view'"))
        for path in ['/api/catalog-import/'+created['id']+'/source','/api/catalog-import/'+created['id']+'/export','/api/business-files/'+str(file)]:
            assert bm.get(path).status_code==403
            assert tsd.get(path).status_code==403
        assert specialist.get('/api/catalog-import/'+created['id']+'/source').status_code==200
        with engine.begin() as db:db.execute(text("UPDATE group_grants SET scope='none' WHERE group_slug='specialist' AND permission_id='nk.export'"))
        assert specialist.get('/api/catalog-import/'+created['id']+'/source').status_code==403
        assert specialist.get('/api/business-files/'+str(file)).status_code==403
        assert specialist.post('/api/catalog-import/preview',params={'organization_id':str(env[2])},content=source_book()).status_code==403


def test_changed_card_identity_is_flagged_instead_of_filling_wrong_pn(env):
    with client(env) as c:
        login(c);created=preview(c,env);result,_=commit(c,created);id=result['card_ids'][0]
        changed=c.put('/api/catalog/'+id,json={**card(env,pn='A-DIFFERENT-PRODUCT',revision=1),'request_id':None})
        assert changed.status_code==200,changed.text
        response=c.get('/api/catalog-import/'+created['id']+'/export');assert response.status_code==200,response.text
        sheet=load_workbook(BytesIO(response.content)).active
        assert sheet['A2'].value=='00123' and sheet['F2'].value in ('',None) and 'требуется сверка' in sheet['G2'].value


def test_commit_replay_after_progress_never_creates_more_cards_and_changed_request_conflicts(env,engine):
    with client(env) as c:
        login(c);created=preview(c,env);result,body=commit(c,created)
        with engine.begin() as db:db.execute(text("UPDATE business_items SET state='processing',revision=revision+3 WHERE id=:id"),{'id':created['id']})
        assert c.post('/api/catalog-import/'+created['id']+'/commit',json=body).json()==result
        assert c.post('/api/catalog-import/'+created['id']+'/commit',json={**body,'reason':'A different request'}).status_code==409
        with engine.begin() as db:db.execute(text("UPDATE group_grants SET scope='none' WHERE group_slug='bm' AND permission_id='nk.import.commit'"))
        assert c.post('/api/catalog-import/'+created['id']+'/commit',json=body).status_code==403
    with engine.connect() as db:assert db.execute(text("SELECT count(*) FROM business_items WHERE kind='card'")).scalar_one()==2


def test_named_sheet_and_colliding_result_headers_do_not_overwrite_input(env):
    book=Workbook();book.active.title='Обложка';book.active['A1']='Read me';sheet=book.create_sheet('Заказ')
    sheet.append(['PN','ТН ВЭД','GTIN ERP','Результат ERP']);sheet.append(['PN/Case','8542399000','Source value','Source status'])
    stream=BytesIO();book.save(stream)
    with client(env) as c:
        login(c)
        assert c.post('/api/catalog-import/preview',params={'organization_id':str(env[1])},content=stream.getvalue()).status_code==422
        created=preview(c,env,stream.getvalue(),sheet_name='Заказ');commit(c,created)
        response=c.get('/api/catalog-import/'+created['id']+'/export');sheet=load_workbook(BytesIO(response.content))['Заказ']
        assert sheet['C2'].value=='Source value' and sheet['D2'].value=='Source status'
        assert sheet['E1'].value=='GTIN ERP (2)' and sheet['F1'].value=='Результат ERP (2)'


def test_active_workbooks_and_unsafe_package_structures_are_rejected_before_persistence(env,engine):
    original=source_book();base=zipfile.ZipFile(BytesIO(original));variants=[]
    def altered(name,data):
        out=BytesIO()
        with zipfile.ZipFile(out,'w') as z:
            for info in base.infolist():z.writestr(info,data if info.filename==name else base.read(info))
            if name not in base.namelist():z.writestr(name,data)
        return out.getvalue()
    variants.append(altered('xl/vbaProject.bin',b'not a macro'))
    variants.append(altered('../outside.txt',b'never extracted'))
    variants.append(altered('xl/workbook.xml',b'<!DOCTYPE workbook [<!ENTITY x "unsafe">]>'+base.read('xl/workbook.xml')))
    book=load_workbook(BytesIO(original));book.active['A2']='=HYPERLINK("file:///private", "PN")';out=BytesIO();book.save(out);variants.append(out.getvalue())
    book=load_workbook(BytesIO(original));book.active['A2'].hyperlink='file:///private';out=BytesIO();book.save(out);variants.append(out.getvalue())
    with client(env) as c:
        login(c)
        for content in variants:
            result=c.post('/api/catalog-import/preview',params={'organization_id':str(env[1])},content=content)
            assert result.status_code==422,result.text
    with engine.connect() as db:
        assert db.execute(text('SELECT count(*) FROM business_items')).scalar_one()==0
        assert db.execute(text('SELECT count(*) FROM business_artifacts')).scalar_one()==0


def test_legacy_import_still_exports_without_inventing_a_source_workbook(env,engine):
    from marking_erp import auth
    from marking_erp.business import store
    with client(env) as c:
        login(c);created=c.post('/api/catalog',json=card(env)).json()
        with engine.begin() as db:
            user=dict(db.execute(text('SELECT * FROM users WHERE id=:id'),{'id':env[3]['bm']}).mappings().one())
            current=auth.Actor(db,user,'fixture',None)
            old=store.create(current,'import','Legacy import',{'card_ids':[created['id']]},env[1],state='done')
        result=c.get('/api/catalog-import/'+str(old['id'])+'/export');assert result.status_code==200,result.text
        rows=list(load_workbook(BytesIO(result.content)).active.values)
        assert len(rows[0])==5 and rows[1][0]=='PN-100'
        assert c.get('/api/catalog-import/'+str(old['id'])+'/source').status_code==404
