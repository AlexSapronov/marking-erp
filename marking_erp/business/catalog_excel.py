"""Bounded XLSX input and a two-column OOXML patch; untouched ZIP members stay exact."""
from io import BytesIO
import posixpath,re,zipfile
from xml.dom import minidom
from urllib.parse import urlsplit
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAIN='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
MIME='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
ALIASES={'pn':('pn','артикул','модель'),'gtin':('gtin',),'tnved':('тнвэд','тн вэд'),
    'name':('название','наименование'),'brand':('товарный знак','бренд'),'category_id':('id категории','category_id')}


def xml(data):
    # Accept UTF-8 OOXML only. Reject DTD/entity declarations before the DOM parser.
    value=data.decode('utf-8-sig')
    if '\x00' in value or '<!DOCTYPE' in value.upper() or '<!ENTITY' in value.upper():
        raise ValueError('Недопустимые объявления XML в Excel')
    return minidom.parseString(value.encode('utf-8'))


def package(content):
    if len(content)>10*1024*1024:raise ValueError('Файл больше 10 МБ')
    archive=zipfile.ZipFile(BytesIO(content));infos=archive.infolist();names=[i.filename for i in infos]
    if len(names)>1500 or len(set(names))!=len(names) or sum(i.file_size for i in infos)>40*1024*1024:
        archive.close();raise ValueError('Слишком большой или неоднозначный XLSX')
    try:
        for info in infos:
            name=info.filename
            if name.startswith('/') or '\\' in name or '..' in name.split('/') or info.flag_bits&1:
                raise ValueError('Недопустимая структура XLSX')
            lowered=name.lower()
            if any(part in lowered for part in ('vbaproject','activex/','embeddings/','externallinks/','_xmlsignatures/')):
                raise ValueError('Для импорта нужна книга XLSX без макросов, вложений и внешних подключений')
            if lowered.endswith(('.xml','.rels')):
                if info.file_size>8*1024*1024:raise ValueError('Слишком большой XML-лист Excel')
                doc=xml(archive.read(info))
                try:
                    if doc.getElementsByTagNameNS(MAIN,'f'):
                        raise ValueError('Сохраните Excel как значения: формулы в партии импорта не поддерживаются')
                    for rel in doc.getElementsByTagNameNS('*','Relationship'):
                        if rel.getAttribute('TargetMode').lower()=='external' and not (rel.getAttribute('Type').endswith('/hyperlink') and urlsplit(rel.getAttribute('Target')).scheme.lower() in ('https','http','mailto')):
                            raise ValueError('Удалите внешние ссылки из книги Excel перед импортом')
                finally:doc.unlink()
        return archive
    except Exception:
        archive.close();raise


def children(node,name):
    return [x for x in node.childNodes if x.nodeType==x.ELEMENT_NODE and x.namespaceURI==MAIN and x.localName==name]


def worksheet_path(archive,name):
    book=xml(archive.read('xl/workbook.xml'));rels=xml(archive.read('xl/_rels/workbook.xml.rels'))
    try:
        matches=[s for s in book.getElementsByTagNameNS(MAIN,'sheet') if s.getAttribute('name')==name]
        if len(matches)!=1:raise ValueError('Не удалось однозначно определить лист Excel')
        rid=matches[0].getAttributeNS(REL,'id')
        matches=[r for r in rels.getElementsByTagNameNS('*','Relationship') if r.getAttribute('Id')==rid]
        if len(matches)!=1:raise ValueError('Не удалось определить содержимое листа')
        target=matches[0].getAttribute('Target')
        path=target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
        if not path.startswith('xl/') or path not in archive.namelist():raise ValueError('Недопустимый путь листа')
        return path
    finally:book.unlink();rels.unlink()


def read(content,sheet_name=None):
    with package(content):pass
    book=load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
    try:
        if len(book.worksheets)>20:raise ValueError('В книге больше 20 листов')
        if sheet_name is not None and sheet_name not in book.sheetnames:raise ValueError('Указанный лист не найден')
        sheet=book[sheet_name] if sheet_name is not None else book.active
        if sheet not in book.worksheets:raise ValueError('Выберите обычный лист с таблицей')
        if sheet.max_row>2001 or sheet.max_column>200:raise ValueError('До 2000 строк данных и 200 столбцов на выбранном листе')
        iterator=sheet.iter_rows();header=next(iterator,());mapping={}
        for key,aliases in ALIASES.items():
            found=[i for i,cell in enumerate(header) if str(cell.value or '').strip().casefold() in aliases]
            if len(found)>1:raise ValueError('Неоднозначные столбцы: '+key)
            mapping[key]=found[0] if found else None
        if mapping['pn'] is None or mapping['tnved'] is None:raise ValueError('Шапка Excel должна содержать PN и ТН ВЭД')
        rows=[]
        for n,row in enumerate(iterator,2):
            if not any(c.value is not None for c in row):continue
            values={k:row[i].value if i is not None and i<len(row) else '' for k,i in mapping.items()}
            error=''
            if not isinstance(values['pn'],str):error='PN должен быть текстом Excel: числовой артикул может потерять ведущие нули'
            for key,value in values.items():
                if isinstance(value,bool):values[key]=str(value)
                elif isinstance(value,(int,float)):
                    try:values[key]=str(int(value)) if value==int(value) else str(value)
                    except (ValueError,OverflowError):values[key]=''
                else:values[key]=str(value if value is not None else '').strip()
            values['gtin']=values['gtin'].zfill(14) if values['gtin'] else ''
            values['tnved']=values['tnved'].zfill(10) if values['tnved'] else ''
            values['name']=values['name'] or values['pn']
            rows.append({'row':n,'values':values,'error':error})
        if not rows:raise ValueError('На выбранном листе нет строк товаров')
        return {'sheet_name':sheet.title,'columns':sheet.max_column,'headers':[str(c.value or '') for c in header],
                'rows':rows,'sheet_names':book.sheetnames}
    finally:book.close()


def append_results(content,meta,results):
    """Append inline strings without round-tripping formatting, drawings or other sheets."""
    with package(content) as archive:
        path=worksheet_path(archive,meta['sheet_name']);doc=xml(archive.read(path))
        try:
            root=doc.documentElement
            if root.namespaceURI!=MAIN:raise ValueError('Неподдерживаемый формат листа')
            data=children(root,'sheetData')
            if len(data)!=1:raise ValueError('Не удалось определить таблицу листа')
            data=data[0];rows={int(r.getAttribute('r')):r for r in children(data,'row')}
            prefix=(root.prefix+':') if root.prefix else ''
            def element(name):return doc.createElementNS(MAIN,prefix+name)
            used=set(meta.get('headers',[]));titles=[]
            for base in ('GTIN ERP','Результат ERP'):
                title=base;n=2
                while title in used:title=f'{base} ({n})';n+=1
                used.add(title);titles.append(title)
            values={1:titles,**results}
            for number,pair in sorted(values.items()):
                row=rows.get(number)
                if row is None:
                    row=element('row');row.setAttribute('r',str(number))
                    later=next((r for r in children(data,'row') if int(r.getAttribute('r'))>number),None)
                    data.insertBefore(row,later);rows[number]=row
                for offset,value in enumerate(pair,1):
                    cell=element('c');cell.setAttribute('r',get_column_letter(meta['columns']+offset)+str(number));cell.setAttribute('t','inlineStr')
                    inline=element('is');text=element('t');text.setAttributeNS('http://www.w3.org/XML/1998/namespace','xml:space','preserve')
                    text.appendChild(doc.createTextNode(re.sub('[\x00-\x08\x0b\x0c\x0e-\x1f]','',str(value))));inline.appendChild(text);cell.appendChild(inline);row.appendChild(cell)
                row.removeAttribute('spans') if row.hasAttribute('spans') else None
            columns=children(root,'cols')
            if columns:columns=columns[0]
            else:columns=element('cols');root.insertBefore(columns,data)
            first,last=meta['columns']+1,meta['columns']+2
            for col in list(children(columns,'col')):
                lo,hi=int(col.getAttribute('min')),int(col.getAttribute('max'))
                if hi<first or lo>last:continue
                for start,end in ((lo,min(hi,first-1)),(max(lo,last+1),hi)):
                    if start<=end:
                        clone=col.cloneNode(True);clone.setAttribute('min',str(start));clone.setAttribute('max',str(end));columns.insertBefore(clone,col)
                columns.removeChild(col)
            for offset,width in ((1,'20'),(2,'64')):
                col=element('col');col.setAttribute('min',str(meta['columns']+offset));col.setAttribute('max',str(meta['columns']+offset));col.setAttribute('width',width);col.setAttribute('customWidth','1');columns.appendChild(col)
            for col in sorted(children(columns,'col'),key=lambda c:int(c.getAttribute('min'))):columns.appendChild(col)
            dimensions=children(root,'dimension')
            if dimensions:
                existing=dimensions[0].getAttribute('ref').split(':')[0]
                dimensions[0].setAttribute('ref',existing+':'+get_column_letter(meta['columns']+2)+str(max(rows)))
            patched=doc.toxml(encoding='utf-8');out=BytesIO()
            with zipfile.ZipFile(out,'w') as target:
                target.comment=archive.comment
                for info in archive.infolist():target.writestr(info,patched if info.filename==path else archive.read(info))
            return out.getvalue()
        finally:doc.unlink()
