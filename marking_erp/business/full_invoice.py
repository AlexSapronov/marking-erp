"""Read a supplier's full invoice and prove emitted marked quantities unchanged."""
from collections import Counter
from io import BytesIO
import zipfile
import openpyxl
from .. import invoice_excel as x

FIELDS={'pn':'PN','gtin':'GTIN','quantity':'Количество товара'}
ALIASES={key:{x.normalized_header(v) for v in values} for key,values in {
    'pn':['PN','Артикул','Part number','P/N','Модель'],
    'gtin':['GTIN','Код товара GTIN','Код товара'],
    'quantity':['Количество','Кол-во','Количество товара','Всего штук','Количество всего','QTY','Quantity','Штук']}.items()}


def parse(content,expected,*,sheet_name=None,header_row=None,mapping=None):
    if not content or len(content)>10*1024*1024:raise ValueError('Нужен XLSX до 10 МБ')
    with zipfile.ZipFile(BytesIO(content)) as archive:
        if sum(i.file_size for i in archive.infolist())>80*1024*1024:raise ValueError('Слишком большой распакованный Excel')
    book=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=False)
    try:
        if len(book.worksheets)>20:raise ValueError('В книге больше 20 листов')
        descriptions=[]
        for ws in book:
            if ws.max_column>200 or ws.max_row>20030:raise ValueError('Допускается до 200 столбцов и 20000 строк')
            best=(0,0,{})
            for n,raw in enumerate(ws.iter_rows(max_row=min(ws.max_row,30)),1):
                found={k:[i for i,cell in enumerate(raw) if x.normalized_header(cell.value) in names] for k,names in ALIASES.items()}
                result={k:v[0] for k,v in found.items() if len(v)==1}
                if len(result)>best[0]:best=(len(result),n,result)
            descriptions.append({'name':ws.title,'header_row':best[1],'mapping':best[2],'columns':[{'index':i,'letter':openpyxl.utils.get_column_letter(i+1)} for i in range(ws.max_column)]})
        matches=[d for d in descriptions if len(d['mapping'])==3]
        chosen=next((d for d in descriptions if d['name']==sheet_name),None) if sheet_name else (matches[0] if len(matches)==1 else None)
        if sheet_name and chosen is None:raise ValueError('Лист не найден')
        if chosen is None or (mapping is None and len(chosen['mapping'])!=3):
            return {'requires_mapping':True,'fields':FIELDS,'sheets':descriptions,'lines':[],'errors':[],'warnings':[]}
        ws=book[chosen['name']];mapping=chosen['mapping'] if mapping is None else mapping
        if set(mapping)!=set(FIELDS) or len(set(mapping.values()))!=3 or any(type(v) is not int or not 0<=v<ws.max_column for v in mapping.values()):raise ValueError('Выберите отдельный столбец PN, GTIN и общего количества')
        header_row=chosen['header_row'] if header_row is None else header_row
        if type(header_row) is not int or not 0<=header_row<=ws.max_row or ws.max_row-header_row>20000:raise ValueError('Проверьте строку заголовков и число строк данных')
        numbers=x._raw_numbers(content,ws.title,set(mapping.values()),header_row+1)
        expected_counts=Counter()
        for line in expected:
            if line.get('marked',True):expected_counts[(line['pn'],line['gtin'])]+=line['quantity']*line['labels']
        marked_pns={p for p,g in expected_counts};actual=Counter();info=[];errors=[];warnings=[]
        for n,row in enumerate(ws.iter_rows(min_row=header_row+1,max_col=ws.max_column),header_row+1):
            selected=[row[mapping[k]] for k in FIELDS]
            if all(c.value is None or str(c.value).strip()=='' for c in selected):continue
            try:
                if any(c.data_type=='f' or (c.data_type=='e' and c.coordinate not in numbers) for c in selected):raise ValueError('Формула или ошибка Excel: вставьте значения')
                raw=[numbers.get(getattr(c,'coordinate',''),c.value) for c in selected]
                if isinstance(raw[0],(int,float)):warnings.append(f'Строка {n}: числовой PN; проверьте ведущие нули')
                pn=x.pn_text(raw[0]);quantity=x.integer_value(raw[2],'Общее количество товара',9999999900000)
                if raw[1] is None or str(raw[1]).strip()=='':
                    if pn in marked_pns:raise ValueError('Этот PN уже маркировался; укажите его GTIN, не переводите его в немаркируемые')
                    info.append({'pn':pn,'gtin':'','tnved':'','quantity':quantity,'labels':0,'marked':False})
                else:
                    gtin=x.identifier(raw[1],14,'GTIN')[0]
                    if not x.valid_gtin(gtin) or (pn,gtin) not in expected_counts:raise ValueError('PN/GTIN отсутствует в уже промаркированном составе')
                    actual[(pn,gtin)]+=quantity
            except ValueError as error:errors.append({'row':n,'message':str(error)})
        for (pn,gtin),quantity in expected_counts.items():
            if actual[(pn,gtin)]!=quantity:errors.append({'row':0,'message':f'{pn} / {gtin}: было промаркировано {quantity}, в полном Excel {actual[(pn,gtin)]}. Изменять маркированную часть нельзя.'})
        if sum(r.get('marked',True) for r in expected)+len(info)>20000:errors.append({'row':0,'message':'Полный состав превышает 20000 строк'})
        return {'requires_mapping':False,'sheet':ws.title,'header_row':header_row,'mapping':mapping,'lines':info,'marked_positions':len(expected_counts),'errors':errors,'warnings':warnings}
    finally:book.close()
