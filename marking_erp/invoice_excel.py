"""Infer Excel fields, request explicit mapping on ambiguity, normalize field values."""

from decimal import Decimal, InvalidOperation
from io import BytesIO
import math
import posixpath
import xml.etree.ElementTree as ET
import re
import unicodedata
import zipfile
import openpyxl
import xlsxwriter
from openpyxl.utils import get_column_letter

HEADERS = ["PN", "GTIN", "Кол-во в 1 этикетке", "Кол-во этикеток", "Код ТН ВЭД"]
FIELDS = dict(zip(("pn", "gtin", "quantity", "labels", "tnved"), HEADERS))
MAX_ROWS, MAX_COLUMNS, HEADER_SCAN_ROWS = 20000, 200, 30
ALIASES = {
    "pn": [
        "PN",
        "Part number",
        "Part no",
        "P/N",
        "Артикул",
        "Модель",
        "Партномер",
        "Model",
        "Article",
        "SKU",
        "Модель/артикул",
    ],
    "gtin": ["GTIN", "ГТИН", "EAN", "EAN13", "EAN14", "Код товара"],
    "quantity": [
        "Кол-во в 1 этикетке",
        "Количество в 1 этикетке",
        "Количество в этикетке",
        "Количество в одной этикетке",
        "Кол-во в одной этикетке",
        "Штук в 1 этикетке",
        "Штук в этикетке",
        "Количество в упаковке",
        "Штук в упаковке",
        "Количество штук в одной этикетке",
        "Qty per label",
        "Qty per pack",
        "Pieces per label",
    ],
    "labels": [
        "Кол-во этикеток",
        "Количество этикеток",
        "Этикеток",
        "Количество упаковок",
        "Кол-во упаковок",
        "Количество КМ",
        "Кол-во КМ",
        "Labels",
        "Label count",
    ],
    "tnved": [
        "Код ТН ВЭД",
        "ТН ВЭД",
        "Код ТНВЭД",
        "ТНВЭД",
        "TN VED",
        "TN VED code",
        "HS code",
        "Код ТН ВЭД ЕАЭС",
    ],
}


def normalized_header(value):
    return re.sub(
        r"[^a-zа-я0-9]",
        "",
        unicodedata.normalize("NFKC", str(value or "")).casefold().replace("ё", "е"),
    )


ALIAS_INDEX = {
    normalized_header(alias): field
    for field, aliases in ALIASES.items()
    for alias in aliases
}


def infer_mapping(values):
    found = {field: [] for field in FIELDS}
    for index, value in enumerate(values):
        field = ALIAS_INDEX.get(normalized_header(value))
        if field:
            found[field].append(index)
    return {
        field: indices[0] if len(indices) == 1 else None
        for field, indices in found.items()
    }


def valid_gtin(s):
    return (
        bool(re.fullmatch("[0-9]{14}", s))
        and int(s[-1])
        == (
            10
            - sum(int(c) * (3 if i % 2 == 0 else 1) for i, c in enumerate(s[:13])) % 10
        )
        % 10
    )


def integer_value(value, label, maximum, minimum=1):
    """Accept stored numbers and numeric text, without rounding or evaluating formulas."""
    if isinstance(value, bool) or value is None:
        raise ValueError(f"{label}: целое число от {minimum} до {maximum}")
    if isinstance(value, (int, float, Decimal)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError(f"{label}: требуется конечное число")
        text = str(value)
    elif isinstance(value, str):
        text = value.strip()
        # Only conventional groups of three, never arbitrary embedded whitespace.
        if re.fullmatch(r"[+]?[0-9]{1,3}(?:[ \u00a0\u202f][0-9]{3})+(?:[.,]0+)?", text):
            text = re.sub(r"[ \u00a0\u202f]", "", text)
        if len(text) > 128 or not re.fullmatch(
            r"[+-]?[0-9]+(?:[.,][0-9]+)?(?:[eE][+-]?[0-9]{1,6})?", text
        ):
            raise ValueError(f"{label}: целое число от {minimum} до {maximum}")
    else:
        raise ValueError(f"{label}: ожидается число или текст с цифрами")
    try:
        number = Decimal(text.replace(",", "."))
        if (
            not number.is_finite()
            or not minimum <= number <= maximum
            or number != number.to_integral_value()
        ):
            raise ValueError(
                f"{label}: целое число от {minimum} до {maximum}, без округления"
            )
        return int(number)
    except InvalidOperation as exc:
        raise ValueError(f"{label}: некорректное число") from exc


def identifier(value, length, label):
    """Fixed-length identifiers can recover zeros lost by numeric Excel cells."""
    if isinstance(value, str) and re.fullmatch("[0-9]+", value.strip()):
        text = value.strip()
        if len(text) > length:
            raise ValueError(f"{label}: не более {length} цифр")
    else:
        text = str(integer_value(value, label, 10**length - 1, minimum=0))
    return text.zfill(length), len(text) < length


def pn_text(value):
    if isinstance(value, bool) or value is None:
        raise ValueError("PN: заполните артикул")
    if isinstance(value, (int, float, Decimal)):
        try:
            number = Decimal(str(value))
            if not number.is_finite() or abs(number) >= 10**15:
                raise ValueError(
                    "PN: числовой артикул слишком длинный; укажите его текстом без потери точности"
                )
            value = (
                format(number, "f").rstrip("0").rstrip(".")
                if number % 1
                else str(int(number))
            )
        except InvalidOperation as exc:
            raise ValueError("PN: некорректное значение") from exc
    if (
        not isinstance(value, str)
        or not value.strip()
        or len(value.strip()) > 120
        or any(ord(c) < 32 for c in value)
    ):
        raise ValueError("PN: нужен артикул от 1 до 120 символов")
    return value.strip()


def validate_rows(rows, row_numbers=None, merge=True):
    errors, warnings, merged, pns, gtins = [], [], {}, {}, {}
    for index, row in enumerate(rows):
        n = row_numbers[index] if row_numbers is not None else index + 2
        try:
            if not isinstance(row, (list, tuple)) or len(row) != 5:
                raise ValueError("Нужно пять полей")
            pn, gtin, quantity, labels, tnved = row
            if isinstance(pn,(int,float,Decimal)) and not isinstance(pn,bool):
                warnings.append(f'Строка {n}: PN сохранён в Excel как число; ведущие нули могли быть потеряны. Сверьте артикул с исходным документом и используйте текстовый формат PN.')
            pn = pn_text(pn)
            gtin, gtin_padded = identifier(gtin, 14, "GTIN")
            if not valid_gtin(gtin):
                raise ValueError("GTIN: неверная контрольная цифра")
            tnved, tnved_padded = identifier(tnved, 10, "ТН ВЭД")
            quantity = integer_value(quantity, "Штук в этикетке", 99999999)
            labels = integer_value(labels, "Этикеток", 100000)
            if pn in pns and pns[pn] != gtin:
                raise ValueError("Один PN соответствует разным GTIN")
            if gtin in gtins and gtins[gtin] != tnved:
                raise ValueError("Один GTIN соответствует разным ТН ВЭД")
            pns[pn], gtins[gtin] = gtin, tnved
            key = (pn, gtin, quantity, tnved) + (() if merge else (index,))
            merged[key] = merged.get(key, 0) + labels
            if gtin_padded:
                warnings.append(
                    f"Строка {n}: GTIN дополнен ведущими нулями до 14 цифр: {gtin}"
                )
            if tnved_padded:
                warnings.append(
                    f"Строка {n}: ТН ВЭД дополнен ведущими нулями до 10 цифр: {tnved}"
                )
            if len(pn) > 35:
                warnings.append(
                    f"Строка {n}: длинный PN; проверьте размещение на этикетке"
                )
        except ValueError as exc:
            errors.append({"row": n, "message": str(exc)})
    if not merged and not errors:
        errors.append({"row": 0, "message": "Нет заполненных строк"})
    if sum(merged.values()) > 100000:
        errors.append({"row": 0, "message": "Больше 100000 этикеток в одном документе"})
    return {
        "lines": [
            dict(pn=k[0], gtin=k[1], quantity=k[2], tnved=k[3], labels=v)
            for k, v in merged.items()
        ],
        "errors": errors,
        "warnings": list(dict.fromkeys(warnings)),
    }


def _display(value):
    if value is None:
        return ""
    return str(value)[:160]


def _sheet_description(ws):
    preview = []
    for n, row in enumerate(
        ws.iter_rows(
            min_row=1, max_row=min(ws.max_row, HEADER_SCAN_ROWS), max_col=ws.max_column
        ),
        1,
    ):
        values = [_display(c.value) for c in row]
        preview.append({"row": n, "values": values, "mapping": infer_mapping(values)})
    ranked = [
        (
            sum(i is not None for i in infer_mapping(r["values"]).values()),
            r["row"],
            infer_mapping(r["values"]),
        )
        for r in preview
    ]
    score, header_row, mapping = (
        max(ranked, key=lambda r: (r[0], -r[1]))
        if ranked
        else (0, 1, {key: None for key in FIELDS})
    )
    return {
        "name": ws.title,
        "row_count": ws.max_row,
        "column_count": ws.max_column,
        "columns": [
            {"index": i, "letter": get_column_letter(i + 1)}
            for i in range(ws.max_column)
        ],
        "suggested_header_row": header_row,
        "suggested_mapping": mapping,
        "matched_fields": score,
        "preview": preview,
    }


def _raw_numbers(content, sheet_name, columns, start_row):
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    prefix = "{" + ns["s"] + "}"
    values = {}
    with zipfile.ZipFile(BytesIO(content)) as archive:
        rels = {
            r.attrib["Id"]: r.attrib["Target"]
            for r in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        }
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        sheet = next(
            s
            for s in workbook.findall("s:sheets/s:sheet", ns)
            if s.attrib["name"] == sheet_name
        )
        target = rels[
            sheet.attrib[
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            ]
        ]
        path = (
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.normpath("xl/" + target)
        )
        letters = {get_column_letter(i + 1) for i in columns}
        with archive.open(path) as stream:
            for _, node in ET.iterparse(stream, events=("end",)):
                if node.tag == prefix + "c":
                    coord = node.get("r", "")
                    match = re.fullmatch(r"([A-Z]+)([0-9]+)", coord)
                    if (
                        match
                        and match[1] in letters
                        and int(match[2]) >= start_row
                        and node.get("t", "n") == "n"
                        and node.find(prefix + "f") is None
                    ):
                        value = node.find(prefix + "v")
                        if value is not None and value.text:
                            try:
                                values[coord] = Decimal(value.text)
                            except InvalidOperation:
                                pass
                    node.clear()
                elif node.tag == prefix + "row":
                    node.clear()
    return values


def import_file(content, mapping=None, sheet_name=None, header_row=None, merge=True):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("Размер Excel превышает 10 МБ")
    book = None
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError("Слишком большой распакованный Excel")
        book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
        if len(book.worksheets) > 20:
            raise ValueError("В книге больше 20 листов")
        for ws in book:
            if ws.max_column > MAX_COLUMNS or ws.max_row > MAX_ROWS + HEADER_SCAN_ROWS:
                raise ValueError(
                    f"Лист «{ws.title}»: допускается до {MAX_COLUMNS} столбцов и {MAX_ROWS} строк данных"
                )
        descriptions = [
            _sheet_description(ws)
            for ws in book
            if ws.max_row > 1 or ws.cell(1, 1).value is not None
        ]
        if not descriptions:
            raise ValueError("В книге нет заполненных листов")
        if sheet_name is not None and sheet_name not in {
            d["name"] for d in descriptions
        }:
            raise ValueError("Выбранный лист не найден")
        complete = [d for d in descriptions if d["matched_fields"] == len(FIELDS)]
        preferred = next((d for d in descriptions if d["name"] == sheet_name), None)
        if preferred is None:
            preferred = (
                complete[0]
                if len(complete) == 1
                else max(descriptions, key=lambda d: d["matched_fields"])
            )
        if mapping is None and (
            len(complete) != 1 or (sheet_name is not None and preferred not in complete)
        ):
            return {
                "requires_mapping": True,
                "fields": FIELDS,
                "sheets": descriptions,
                "selected_sheet": preferred["name"],
                "lines": [],
                "errors": [],
                "warnings": [],
            }
        ws = book[preferred["name"]]
        if mapping is None:
            mapping = preferred["suggested_mapping"]
        if (
            not isinstance(mapping, dict)
            or set(mapping) != set(FIELDS)
            or any(
                type(i) is not int or not 0 <= i < ws.max_column
                for i in mapping.values()
            )
            or len(set(mapping.values())) != len(FIELDS)
        ):
            raise ValueError("Укажите отдельный столбец для каждого из пяти полей")
        header_row = (
            preferred["suggested_header_row"] if header_row is None else header_row
        )
        if type(header_row) is not int or not 0 <= header_row <= ws.max_row:
            raise ValueError(
                "Проверьте номер строки заголовков; 0 означает файл без заголовков"
            )
        data, physical_rows, cell_errors = [], [], []
        numeric_values = _raw_numbers(
            content, ws.title, set(mapping.values()), header_row + 1
        )
        if ws.max_row - header_row > MAX_ROWS:
            raise ValueError(f"В выбранном листе больше {MAX_ROWS} строк данных")
        for n, row in enumerate(
            ws.iter_rows(min_row=header_row + 1, max_col=ws.max_column), header_row + 1
        ):
            selected = [row[mapping[field]] for field in FIELDS]
            if all(
                c.value is None or (isinstance(c.value, str) and not c.value.strip())
                for c in selected
            ):
                continue
            invalid = [
                f"{FIELDS[field]} ({get_column_letter(mapping[field] + 1)})"
                for field, cell in zip(FIELDS, selected)
                if cell.data_type == "f"
                or (cell.data_type == "e" and cell.coordinate not in numeric_values)
            ]
            if invalid:
                cell_errors.append(
                    {
                        "row": n,
                        "message": "Формула или ошибка Excel в поле "
                        + ", ".join(invalid)
                        + ". Вставьте рассчитанные значения.",
                    }
                )
                continue
            data.append([numeric_values.get(c.coordinate, c.value) for c in selected])
            physical_rows.append(n)
        result = validate_rows(data, row_numbers=physical_rows, merge=merge)
        if cell_errors:
            result["errors"] = sorted(
                cell_errors + [e for e in result["errors"] if e["row"]],
                key=lambda e: e["row"],
            )
        result.update(
            requires_mapping=False,
            sheet=ws.title,
            header_row=header_row,
            mapping=mapping,
        )
        return result
    except ValueError:
        raise
    except zipfile.BadZipFile as exc:
        raise ValueError(
            "Не удалось прочитать XLSX-файл. Сохраните книгу в формате .xlsx"
        ) from exc
    except Exception as exc:
        raise ValueError(
            "Не удалось прочитать XLSX-файл. Сохраните книгу в формате .xlsx"
        ) from exc
    finally:
        if book is not None:
            book.close()


def xlsx(sheets):
    buf = BytesIO()
    wb = xlsxwriter.Workbook(
        buf, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False}
    )
    head = wb.add_format({"bold": True, "bg_color": "#D8E9E7", "border": 1})
    text = wb.add_format({"num_format": "@"})
    for name, rows in sheets.items():
        ws = wb.add_worksheet(name)
        ws.freeze_panes(1, 0)
        ws.set_column(0, 4, 28, text)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if isinstance(value, str):
                    ws.write_string(r, c, value, head if r == 0 else text)
                elif value is not None:
                    ws.write_number(r, c, value, head if r == 0 else None)
    wb.close()
    return buf.getvalue()


def template():
    # Text identifiers, numeric quantities; set formats for future user input as well.
    buf = BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    ws = wb.add_worksheet("Товары")
    text = wb.add_format({"num_format": "@"})
    number = wb.add_format({"num_format": "0"})
    head = wb.add_format({"bold": True, "bg_color": "#D8E9E7", "border": 1})
    for c, title in enumerate(HEADERS):
        ws.set_column(c, c, 28, number if c in (2, 3) else text)
        ws.write_string(0, c, title, head)
    ws.freeze_panes(1, 0)
    wb.close()
    return buf.getvalue()
