"""Bounded fallback XLSX: one complete expected KM per marked row, no API calls."""
from io import BytesIO
from hashlib import sha256
import openpyxl
from . import gs1,invoice_excel as excel
from .xlsx_package import package

HEADERS=['PN','GTIN','Код маркировки','Количество в КМ','ТН ВЭД']
MAX_ROWS=10000


def template():
    return excel.xlsx({'Товары':[HEADERS]})


def parse_code(raw):
    value=raw.strip()
    # Prefer real scanner bytes. Printable notation may also be valid data
    # inside the serial/crypto: replace it only at known GS1 boundaries.
    try:return gs1.parse(value.encode('ascii'),scanner=True)
    except ValueError:pass
    value=value[3:] if value.startswith(']d2') else value
    def separator(at):
        nonlocal value
        for marker in ('<GS>','_x001D_','_x001d_'):
            if value[at:].startswith(marker):
                value=value[:at]+'\x1d'+value[at+len(marker):];break
    separator(38)
    if value[39:41]=='91':separator(45);end=92
    else:end=45
    separator(end)
    return gs1.parse(value.encode('ascii'),scanner=True)


def read(content):
    with package(content):pass
    book=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
    try:
        if len(book.worksheets)>20 or 'Товары' not in book.sheetnames:raise ValueError('Используйте шаблон с листом «Товары»')
        sheet=book['Товары']
        if sheet.max_row>MAX_ROWS+1 or sheet.max_column>200:raise ValueError('До 10000 строк данных и 200 столбцов')
        rows=sheet.iter_rows(max_col=5);header=next(rows,())
        if [c.value for c in header[:4]]!=HEADERS[:4] or (len(header)>4 and header[4].value not in (None,'ТН ВЭД','Код ТН ВЭД')):
            raise ValueError('Шапка: PN / GTIN / Код маркировки / Количество в КМ / ТН ВЭД (необязательно)')
        lines=[];indices={};expected=[];evidence={};errors=[];error_count=0;unmarked=0;warnings=[];pns={};categories={}
        for n,row in enumerate(rows,2):
            values=[c.value for c in row]
            if all(v is None or (isinstance(v,str) and not v.strip()) for v in values):continue
            try:
                pn,gtin,raw,quantity,tnved=values
                if not isinstance(pn,str):raise ValueError('PN должен быть текстом Excel: числовой артикул может потерять ведущие нули')
                pn=excel.pn_text(pn);quantity=excel.integer_value(quantity,'Количество в КМ',99999999)
                tnved='' if tnved in (None,'') else excel.identifier(tnved,10,'ТН ВЭД')[0]
                if gtin in (None,''):
                    if raw not in (None,''):raise ValueError('КМ без GTIN: заполните GTIN либо удалите код у немаркируемой строки')
                    lines.append({'pn':pn,'gtin':'','quantity':quantity,'labels':0,'tnved':tnved,'marked':False});unmarked+=1
                    continue
                gtin,padded=excel.identifier(gtin,14,'GTIN')
                if not excel.valid_gtin(gtin):raise ValueError('Неверная контрольная цифра GTIN')
                if not isinstance(raw,str) or len(raw)>2048:raise ValueError('Нужен полный текст КМ с криптохвостом')
                if len(raw.strip())==38:raise ValueError('Нужен полный КМ с криптохвостом, а не только КИ')
                parsed=parse_code(raw)
                if parsed.gtin!=gtin:raise ValueError('GTIN строки не совпадает с КМ')
                if (parsed.quantity or 1)!=quantity:raise ValueError('Количество строки не совпадает с AI30 полного КМ (без AI30 — 1)')
                if parsed.cis in evidence:raise ValueError('Повтор КМ в исходном файле')
                if pn in pns and pns[pn]!=gtin:raise ValueError('Один PN соответствует разным GTIN')
                if gtin in categories and categories[gtin] and tnved and categories[gtin]!=tnved:raise ValueError('Один GTIN соответствует разным ТН ВЭД')
                pns[pn]=gtin;categories[gtin]=tnved or categories.get(gtin,'')
                key=(pn,gtin,quantity,tnved)
                if key not in indices:
                    indices[key]=len(lines);lines.append({'pn':pn,'gtin':gtin,'quantity':quantity,'labels':0,'tnved':tnved})
                index=indices[key];lines[index]['labels']+=1;expected.append(parsed.cis)
                evidence[parsed.cis]={'sha256':sha256(parsed.normalized).hexdigest(),'quantity':quantity,'line_index':index}
                if padded and len(warnings)<50:warnings.append(f'Строка {n}: GTIN дополнен ведущими нулями')
            except (ValueError,UnicodeError) as error:
                error_count+=1
                if len(errors)<100:errors.append({'row':n,'message':str(error) if isinstance(error,ValueError) else 'КМ должен содержать ASCII-символы'})
        if not lines and not errors:errors=[{'row':0,'message':'Нет заполненных строк'}];error_count=1
        return {'lines':lines,'expected_codes':expected,'expected_payloads':evidence,'errors':errors,'error_count':error_count,
            'warnings':list(dict.fromkeys(warnings)),'unmarked_lines':unmarked,'parser':'warehouse-xlsx-v1'}
    finally:book.close()
